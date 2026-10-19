# ============================================================
#  Nexus Attendance System · index.py (2026 · No Group Edition)
# ============================================================
import os
import secrets
import time
import hmac
import hashlib
import io
import csv
import random
import re
import urllib.parse
import base64
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta
from collections import defaultdict
import pytz
from flask import Flask, render_template, request, jsonify, session, Response, redirect, make_response
from pymongo import MongoClient, UpdateOne
from werkzeug.security import generate_password_hash, check_password_hash
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address


BASE_DIR = os.path.abspath(os.path.dirname(__file__))
PARENT_DIR = os.path.abspath(os.path.join(BASE_DIR, '..'))

TEMPLATE_DIR = (
    os.path.join(PARENT_DIR, 'templates')
    if os.path.exists(os.path.join(PARENT_DIR, 'templates'))
    else os.path.join(BASE_DIR, 'templates')
)
STATIC_DIR = (
    os.path.join(PARENT_DIR, 'static')
    if os.path.exists(os.path.join(PARENT_DIR, 'static'))
    else os.path.join(BASE_DIR, 'static')
)

app = Flask(__name__, template_folder=TEMPLATE_DIR, static_folder=STATIC_DIR, static_url_path='/static')

app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0
app.jinja_env.auto_reload = True
app.jinja_env.cache = {}

SECRET_KEY = os.environ.get("SECRET_KEY", "Nexus_Attendance_Super_Key_2026_Fixed_ChangeMe")
app.secret_key = SECRET_KEY
app.permanent_session_lifetime = timedelta(days=7)

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(days=7),
    MAX_CONTENT_LENGTH=16 * 1024 * 1024,
)

limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["3000 per day", "800 per hour"],
    storage_uri="memory://"
)


_db_user = urllib.parse.quote_plus(os.environ.get('MONGO_USER', 'ahmedattia20041120_db_user'))
_db_pass = urllib.parse.quote_plus(os.environ.get('MONGO_PASS', 'wjXYBO8Pbj5GijfS'))
DEFAULT_MONGO_URI = (
    f"mongodb+srv://{_db_user}:{_db_pass}"
    "@cluster0.yimrrnh.mongodb.net/?retryWrites=true&w=majority&appName=Cluster0"
)
MONGO_URI = os.environ.get("MONGO_URI", DEFAULT_MONGO_URI)

try:
    client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=10000, tz_aware=False)
    db = client['nexus_attendance_system']
    client.admin.command('ping')
except Exception as _e:
    raise SystemExit(f"❌ فشل الاتصال بقاعدة البيانات: {_e}")

users_col = db['users']
subjects_col = db['subjects']
sessions_col = db['sessions']
attendance_col = db['attendance']
students_col = db['students']
quizzes_col = db['quizzes']
quiz_attempts_col = db['quiz_attempts']
quiz_blocks_col = db['quiz_blocks']


def _safe_create_index(col, keys, **kwargs):
    try:
        col.create_index(keys, **kwargs)
    except Exception:
        pass

_safe_create_index(users_col, "username", unique=True)
_safe_create_index(students_col, "student_id", unique=True)
_safe_create_index(attendance_col, [("student_id", 1), ("session_id", 1)], unique=True)
_safe_create_index(sessions_col, "session_id", unique=True)
_safe_create_index(sessions_col, "screen_secret", sparse=True)
_safe_create_index(quizzes_col, "id", unique=True)
_safe_create_index(quiz_attempts_col, "attempt_id", unique=True)
_safe_create_index(quiz_attempts_col, [("quiz_id", 1), ("student_id", 1)])
_safe_create_index(quiz_blocks_col, [("quiz_id", 1), ("student_id", 1)], unique=True)

# === Performance Indexes ===
_safe_create_index(sessions_col, [("created_at_dt", -1)])
_safe_create_index(sessions_col, [("is_open", 1), ("created_at_dt", -1)])
_safe_create_index(sessions_col, [("subject_id", 1), ("created_at_dt", -1)])
_safe_create_index(quiz_attempts_col, [("quiz_id", 1), ("status", 1)])
_safe_create_index(quiz_blocks_col, [("quiz_id", 1), ("unblocked_at", 1)])
_safe_create_index(attendance_col, [("session_id", 1)])
_safe_create_index(users_col, [("role", 1)])
_safe_create_index(quizzes_col, [("subject_id", 1), ("created_at_dt", -1)])


ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "Nexus_Admin_Core#2026")
ADMIN_PASSWORD_PLAIN = os.environ.get("ADMIN_PASSWORD", "Nx!99@bATU#xK82_Secured")

_existing_admin = users_col.find_one({"username": ADMIN_USERNAME})
if not _existing_admin:
    users_col.insert_one({
        "username": ADMIN_USERNAME,
        "role": "super_admin",
        "name": "الآدمن الرئيسي",
        "is_active": True,
        "password": generate_password_hash(ADMIN_PASSWORD_PLAIN),
        "created_at": datetime.utcnow()
    })


SECRET_SALT        = b"NEXUS_ATTENDANCE_CORE_SECRET_2026_PROD"
QR_SECRET_SALT     = b"NEXUS_QR_ATTENDANCE_SECRET_2026_PROD"
SCREEN_SALT        = b"NEXUS_SCREEN_PAIR_SALT_2026_PROD"

BASE32_CHARS = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"

MANUAL_STEP_INTERVAL = 10
QR_STEP_INTERVAL     = 3
MANUAL_GRACE_PERIOD  = 2.0

SCREEN_STEP_INTERVAL = 20
SCREEN_GRACE_PERIOD  = 2.0

TA_SESSION_AUTO_CLOSE_MINUTES = 15
HEARTBEAT_GRACE_SEC = 30  # ⬅️ raised from 8 → 30 (transient disconnects tolerated)

_SAFE_ERROR_MSG = "حدث خطأ، حاول لاحقاً"

_screen_pending = {}


# ============================================================
#  ⚡ PERFORMANCE: Throttle auto-close
# ============================================================
_last_ta_close_time = 0

def _auto_close_ta_sessions():
    global _last_ta_close_time
    now = time.time()
    if now - _last_ta_close_time < 60:
        return
    _last_ta_close_time = now
    try:
        cutoff = datetime.utcnow() - timedelta(minutes=TA_SESSION_AUTO_CLOSE_MINUTES)
        sessions_col.update_many(
            {"is_open": True, "creator_role": "ta", "created_at_dt": {"$lt": cutoff}},
            {"$set": {"is_open": False, "auto_closed": True, "closed_at": _cairo_now()}}
        )
    except Exception:
        pass


def _generate_code(secret_salt: bytes, step: int, session_id: str) -> str:
    key = secret_salt + session_id.encode('utf-8')
    digest = hmac.new(key, str(step).encode('utf-8'), hashlib.sha256).digest()
    num = int.from_bytes(digest[:5], 'big')
    code = ""
    for _ in range(6):
        code += BASE32_CHARS[num % len(BASE32_CHARS)]
        num //= len(BASE32_CHARS)
    return code


def get_step_code(step: int, session_id: str) -> str:
    return _generate_code(SECRET_SALT, step, session_id)


def get_qr_step_code(step: int, session_id: str) -> str:
    return _generate_code(QR_SECRET_SALT, step, session_id)


def _safe_eq(a: str, b: str) -> bool:
    try:
        return hmac.compare_digest(a.encode('utf-8'), b.encode('utf-8'))
    except Exception:
        return False


def verify_manual_code(user_code: str, session_id: str) -> bool:
    clean = (user_code or '').strip().upper()
    if not clean or len(clean) != 6:
        return False
    now = time.time()
    current_step = int(now // MANUAL_STEP_INTERVAL)
    if _safe_eq(clean, get_step_code(current_step, session_id)):
        return True
    elapsed = now - (current_step * MANUAL_STEP_INTERVAL)
    if elapsed < MANUAL_GRACE_PERIOD and _safe_eq(clean, get_step_code(current_step - 1, session_id)):
        return True
    return False


def verify_qr_code(user_code: str, session_id: str) -> bool:
    clean = (user_code or '').strip().upper()
    if not clean or len(clean) != 6:
        return False
    now = time.time()
    current_step = int(now // QR_STEP_INTERVAL)
    return _safe_eq(clean, get_qr_step_code(current_step, session_id))


def generate_screen_secret() -> str:
    return ''.join(random.choices('0123456789abcdef', k=32))


def get_screen_code(secret: str, step: int) -> str:
    key = SCREEN_SALT + secret.encode('utf-8')
    digest = hmac.new(key, str(step).encode('utf-8'), hashlib.sha256).digest()
    num = int.from_bytes(digest[:8], 'big')
    code = ""
    for _ in range(8):
        code += str(num % 10)
        num //= 10
    return code


def verify_screen_code(user_code: str, secret: str) -> bool:
    clean = (user_code or '').strip()
    if not clean or len(clean) != 8 or not clean.isdigit():
        return False
    now = time.time()
    current_step = int(now // SCREEN_STEP_INTERVAL)
    if _safe_eq(clean, get_screen_code(secret, current_step)):
        return True
    elapsed = now - (current_step * SCREEN_STEP_INTERVAL)
    if elapsed < SCREEN_GRACE_PERIOD and _safe_eq(clean, get_screen_code(secret, current_step - 1)):
        return True
    return False


try:
    for sess in sessions_col.find({"screen_secret": {"$exists": False}}):
        sessions_col.update_one({"_id": sess["_id"]},
            {"$set": {"screen_secret": generate_screen_secret()}})
except Exception:
    pass

try:
    for sess in sessions_col.find({"created_at_dt": {"$exists": False}}):
        sessions_col.update_one({"_id": sess["_id"]},
            {"$set": {"created_at_dt": datetime.utcnow(),
                      "creator_role": sess.get("creator_role", "doctor")}})
except Exception:
    pass


@app.after_request
def set_security_headers(response):
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'SAMEORIGIN'
    response.headers['X-XSS-Protection'] = '1; mode=block'
    response.headers['Referrer-Policy'] = 'strict-origin-when-cross-origin'
    response.headers['Permissions-Policy'] = 'camera=(self), microphone=(), geolocation=()'
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0, private, no-transform'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    return response


@app.errorhandler(429)
def ratelimit_handler(e):
    return jsonify({"status": "error", "message": "تم تجاوز عدد المحاولات! انتظر قليلاً."}), 429


@app.errorhandler(413)
def too_large_handler(e):
    return jsonify({"status": "error", "message": "الملف كبير جداً"}), 413


@app.errorhandler(500)
def internal_error(e):
    try:
        import traceback
        print("=" * 60)
        print("🔥 INTERNAL ERROR:")
        traceback.print_exc()
        print("=" * 60)
    except Exception:
        pass
    return jsonify({"status": "error", "message": _SAFE_ERROR_MSG}), 500


def _clean_str(v, maxlen=200) -> str:
    s = str(v or '').strip()
    return s[:maxlen]


def _is_admin() -> bool:
    return 'admin' in session


def _cairo_now() -> str:
    return datetime.now(pytz.timezone('Africa/Cairo')).strftime("%Y-%m-%d %I:%M:%S %p")


def _client_ip() -> str:
    fwd = request.headers.get('x-forwarded-for', '')
    if fwd:
        return fwd.split(',')[0].strip()
    return request.remote_addr or 'unknown'


@app.route('/')
def student_ui():
    return render_template('index.html')


@app.route('/secure-auth-gateway-2026-x9v2-pl7q-a84m')
def admin_ui():
    return render_template('admin.html')


@app.route('/dis')
def display_screen():
    try:
        app.jinja_env.cache = {}
    except Exception:
        pass
    display_path = os.path.join(TEMPLATE_DIR, 'display.html')
    try:
        with open(display_path, 'r', encoding='utf-8') as f:
            html_content = f.read()
    except Exception:
        html_content = render_template('display.html')
    resp = make_response(html_content)
    resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0, private, no-transform, proxy-revalidate'
    resp.headers['Pragma'] = 'no-cache'
    resp.headers['Expires'] = '0'
    resp.headers['Surrogate-Control'] = 'no-store'
    resp.headers['Vary'] = '*'
    resp.headers['Clear-Site-Data'] = '"cache"'
    resp.headers['Content-Type'] = 'text/html; charset=utf-8'
    return resp


@app.route('/logout-gateway-vip-x9v2-pL7q-2026')
def logout():
    session.clear()
    return redirect('/secure-auth-gateway-2026-x9v2-pl7q-a84m')


@app.route('/api/student-login', methods=['POST'])
@limiter.limit("30 per minute")
def student_login():
    _auto_close_ta_sessions()
    data = request.get_json(force=True, silent=True) or {}
    s_name = _clean_str(data.get('student_name'), 100)
    s_id = _clean_str(data.get('student_id'), 20)
    year = _clean_str(data.get('year'), 40)
    dept = _clean_str(data.get('department', 'عام (IT)'), 60)

    if len(s_name) < 3:
        return jsonify({"status": "error", "message": "يرجى كتابة اسم الطالب بشكل صحيح!"}), 400
    if len(s_id) != 7 or not s_id.isdigit():
        return jsonify({"status": "error", "message": "كود الطالب يجب أن يتكون من 7 أرقام!"}), 400
    if not year:
        return jsonify({"status": "error", "message": "يرجى اختيار الفرقة الدراسية!"}), 400
    if year in ('الفرقة الثالثة', 'الفرقة الرابعة') and not dept:
        return jsonify({"status": "error", "message": "يرجى اختيار القسم!"}), 400

    students_col.update_one({"student_id": s_id},
        {"$set": {"student_id": s_id, "name": s_name,
                  "email": f"{s_id}@batechu.com", "year": year,
                  "department": dept, "last_active": _cairo_now()}},
        upsert=True)

    return jsonify({"status": "success",
        "student": {"student_id": s_id, "name": s_name,
                    "email": f"{s_id}@batechu.com",
                    "year": year, "department": dept}})


@app.route('/api/student-init')
def student_init():
    _auto_close_ta_sessions()
    year = _clean_str(request.args.get('year'), 40)
    dept = _clean_str(request.args.get('dept'), 60)
    query = {}
    if year:
        query["year"] = year
        if dept:
            query["department"] = {"$in": [dept, "عام (IT)"]}
    subs = list(subjects_col.find(query, {"_id": 0}))
    subject_ids = [s['id'] for s in subs]
    all_sessions = list(sessions_col.find({"is_open": True}, {"_id": 0, "screen_secret": 0}))
    filtered = [s for s in all_sessions if s.get('subject_id') in subject_ids]
    return jsonify({"status": "success", "subjects": subs, "sessions": filtered})


@app.route('/api/student-history', methods=['POST'])
@limiter.limit("60 per minute")
def student_history():
    data = request.get_json(force=True, silent=True) or {}
    s_id = _clean_str(data.get('student_id'), 20)
    if len(s_id) != 7 or not s_id.isdigit():
        return jsonify({"status": "error", "message": "ID غير صحيح"}), 400
    records = list(attendance_col.find({"student_id": s_id}, {"_id": 0}).sort("timestamp", -1).limit(200))
    return jsonify({"status": "success", "history": records})


@app.route('/api/submit-attendance', methods=['POST'])
@limiter.limit("40 per minute")
def submit_attendance():
    data = request.get_json(force=True, silent=True) or {}
    s_id = _clean_str(data.get('student_id'), 20)
    s_name = _clean_str(data.get('student_name'), 100)
    session_id = _clean_str(data.get('session_id'), 64)
    code = _clean_str(data.get('code'), 12).upper()
    device_token = _clean_str(data.get('device_token'), 100)
    scan_type = _clean_str(data.get('scan_type', 'manual'), 20).lower()

    if len(s_id) != 7 or not s_id.isdigit():
        return jsonify({"status": "error", "message": "رقم الـ ID غير صحيح!"}), 400

    sess = sessions_col.find_one({"session_id": session_id, "is_open": True})
    if not sess:
        return jsonify({"status": "error", "message": "عفواً، هذه الجلسة مغلقة حالياً أو انتهت!"}), 400

    is_valid = verify_qr_code(code, session_id) if scan_type == 'qr' else verify_manual_code(code, session_id)
    if not is_valid:
        if scan_type == 'qr':
            return jsonify({"status": "error", "message": "كود QR انتهت صلاحيته!"}), 400
        return jsonify({"status": "error", "message": "الرمز السري غير صحيح أو انتهت صلاحيته!"}), 400

    if device_token:
        dup_dev = attendance_col.find_one({"session_id": session_id, "device_token": device_token})
        if dup_dev and dup_dev.get('student_id') != s_id:
            return jsonify({"status": "error", "message": "ممنوع الغش! تم تسجيل حضور طالب آخر من هذا الجهاز."}), 403

    if attendance_col.find_one({"student_id": s_id, "session_id": session_id}):
        return jsonify({"status": "error", "message": "لقد قمت بتسجيل الحضور مسبقاً!"}), 409

    st_record = students_col.find_one({"student_id": s_id})
    year = st_record.get('year', '') if st_record else ''
    dept = st_record.get('department', '') if st_record else ''

    attendance_col.insert_one({
        "student_id": s_id, "student_name": s_name, "year": year, "department": dept,
        "subject_id": sess['subject_id'], "subject_name": sess['subject_name'],
        "session_id": session_id, "session_title": sess['title'],
        "session_type": sess.get('type', 'Lecture'), "timestamp": _cairo_now(),
        "device_token": device_token, "ip": _client_ip(),
        "is_manual": False, "scan_type": scan_type
    })

    return jsonify({"status": "success", "message": f"تم تأكيد حضورك في ({sess['subject_name']} - {sess['title']})"})


@app.route('/api/admin-login', methods=['POST'])
@limiter.limit("15 per minute")
def admin_login():
    data = request.get_json(force=True, silent=True) or {}
    username = _clean_str(data.get('username'), 100)
    password = str(data.get('password', ''))[:300]
    if not username or not password:
        return jsonify({"status": "error", "message": "بيانات ناقصة"}), 400

    user = users_col.find_one({"username": {"$regex": f"^{re.escape(username)}$", "$options": "i"}})
    if user:
        if user.get('is_active', True) is False:
            return jsonify({"status": "error", "message": "هذا الحساب موقوف حالياً."}), 403
        stored_hash = user.get('password', '')
        ok = False
        try:
            ok = check_password_hash(stored_hash, password)
        except Exception:
            ok = False
        if ok:
            session.permanent = True
            session['admin'] = {
                "username": user['username'],
                "name": user.get('name', user['username']),
                "role": user.get('role', 'doctor'),
                "allowed_subjects": user.get('allowed_subjects', [])
            }
            return jsonify({"status": "success", "admin": session['admin']})
    return jsonify({"status": "error", "message": "اسم المستخدم أو كلمة المرور غير صحيحة!"}), 401


@app.route('/api/admin-data')
def get_admin_data():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    _auto_close_ta_sessions()
    role = session['admin']['role']
    curr_username = session['admin']['username']
    curr_admin = users_col.find_one({"username": curr_username})
    allowed_subs = curr_admin.get('allowed_subjects', []) if curr_admin else []
    subjects = []; sessions_list = []; staff_list = []

    if role == 'super_admin':
        subjects = list(subjects_col.find({}, {"_id": 0}))
        sessions_list = list(sessions_col.find(
            {}, {"_id": 0, "screen_secret": 0}
        ).sort("created_at_dt", -1).limit(200))
        staff_list = list(users_col.find({"role": {"$ne": "super_admin"}}, {"_id": 0, "password": 0}))
    elif role == 'doctor':
        subjects = list(subjects_col.find({"id": {"$in": allowed_subs}}, {"_id": 0}))
        sessions_list = list(sessions_col.find(
            {"subject_id": {"$in": allowed_subs}},
            {"_id": 0, "screen_secret": 0}
        ).sort("created_at_dt", -1).limit(200))
        all_tas = list(users_col.find({"role": "ta"}, {"_id": 0, "password": 0}))
        filtered = []
        for ta in all_tas:
            ta_subs = ta.get('allowed_subjects', [])
            has_common = any(s in allowed_subs for s in ta_subs)
            created_by_him = (ta.get('created_by') == curr_username)
            if has_common or created_by_him:
                ta['allowed_subjects'] = [s for s in ta_subs if s in allowed_subs]
                filtered.append(ta)
        staff_list = filtered
    elif role == 'ta':
        subjects = list(subjects_col.find({"id": {"$in": allowed_subs}}, {"_id": 0}))
        sessions_list = list(sessions_col.find({
            "subject_id": {"$in": allowed_subs},
            "type": "Section",
            "created_by_username": curr_username
        }, {"_id": 0, "screen_secret": 0}).sort("created_at_dt", -1).limit(200))
        staff_list = []

    return jsonify({"status": "success",
        "currentAdmin": {"username": session['admin']['username'],
                         "name": session['admin']['name'],
                         "role": session['admin']['role'],
                         "allowed_subjects": allowed_subs},
        "subjects": subjects, "sessions": sessions_list, "staff": staff_list})


@app.route('/api/session-attendance')
def get_session_attendance():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    sess_id = _clean_str(request.args.get('session_id'), 64)
    sess = sessions_col.find_one({"session_id": sess_id}, {"_id": 0, "screen_secret": 0})
    if not sess:
        return jsonify({"status": "error", "message": "الجلسة غير موجودة!"}), 404
    if not _can_access_session(session['admin'], sess_id):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403
    records = list(attendance_col.find({"session_id": sess_id}, {"_id": 0}).sort("timestamp", -1).limit(1000))
    return jsonify({"status": "success", "session": sess, "records": records})


@app.route('/api/live-code')
def get_live_code():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    _auto_close_ta_sessions()
    session_id = _clean_str(request.args.get('session_id'), 64)
    if not session_id:
        return jsonify({"status": "error", "message": "missing"}), 400
    if not _can_access_session(session['admin'], session_id):
        return jsonify({"status": "unauthorized"}), 403

    now = time.time()
    m_step = int(now // MANUAL_STEP_INTERVAL)
    m_remaining = MANUAL_STEP_INTERVAL - (now - (m_step * MANUAL_STEP_INTERVAL))
    manual_code = get_step_code(m_step, session_id)
    q_step = int(now // QR_STEP_INTERVAL)
    q_remaining = QR_STEP_INTERVAL - (now - (q_step * QR_STEP_INTERVAL))
    qr_code = get_qr_step_code(q_step, session_id)

    return jsonify({"manual_code": manual_code, "manual_remaining": m_remaining,
                    "manual_interval": MANUAL_STEP_INTERVAL, "qr_code": qr_code,
                    "qr_remaining": q_remaining, "qr_interval": QR_STEP_INTERVAL,
                    "code": manual_code, "remaining": m_remaining,
                    "interval": MANUAL_STEP_INTERVAL})


def _can_access_session(admin: dict, session_id: str) -> bool:
    sess = sessions_col.find_one({"session_id": session_id}, {"subject_id": 1, "type": 1, "created_by_username": 1})
    if not sess:
        return False
    role = admin.get('role')
    if role == 'super_admin':
        return True
    curr = users_col.find_one({"username": admin.get('username')})
    if not curr:
        return False
    subs = curr.get('allowed_subjects', [])
    if sess.get('subject_id') not in subs:
        return False
    if role == 'doctor':
        return True
    if role == 'ta':
        return (sess.get('type') == 'Section' and sess.get('created_by_username') == curr.get('username'))
    return False


@app.route('/api/session-screen-code')
@limiter.limit("120 per minute")
def session_screen_code():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    session_id = _clean_str(request.args.get('session_id'), 64)
    if not session_id:
        return jsonify({"status": "error", "message": "missing session_id"}), 400
    if not _can_access_session(session['admin'], session_id):
        return jsonify({"status": "unauthorized"}), 403
    sess = sessions_col.find_one({"session_id": session_id})
    if not sess:
        return jsonify({"status": "error", "message": "not found"}), 404
    secret = sess.get('screen_secret')
    if not secret:
        secret = generate_screen_secret()
        sessions_col.update_one({"session_id": session_id}, {"$set": {"screen_secret": secret}})
    now = time.time()
    step = int(now // SCREEN_STEP_INTERVAL)
    remaining = SCREEN_STEP_INTERVAL - (now - (step * SCREEN_STEP_INTERVAL))
    code = get_screen_code(secret, step)
    return jsonify({"status": "success", "code": code, "remaining": remaining, "interval": SCREEN_STEP_INTERVAL})


@app.route('/api/screen-new-token', methods=['POST'])
@limiter.limit("60 per minute")
def screen_new_token():
    token = 'SCR_' + ''.join(random.choices('0123456789abcdef', k=16))
    _screen_pending[token] = {'created': time.time(), 'session_id': None, 'secret': None, 'session_info': None}
    cutoff = time.time() - 600
    for k in list(_screen_pending.keys()):
        if _screen_pending[k].get('created', 0) < cutoff:
            del _screen_pending[k]
    return jsonify({"status": "success", "token": token})


@app.route('/api/screen-check-pair')
@limiter.limit("600 per minute")
def screen_check_pair():
    token = _clean_str(request.args.get('token'), 80)
    if not token or token not in _screen_pending:
        return jsonify({"status": "expired"})
    data = _screen_pending[token]
    if data.get('session_id'):
        return jsonify({"status": "paired", "session_id": data['session_id'],
                        "screen_secret": data['secret'], "session_info": data['session_info']})
    return jsonify({"status": "waiting"})


@app.route('/api/screen-pair-token', methods=['POST'])
@limiter.limit("30 per minute")
def screen_pair_token():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    token = _clean_str(data.get('token'), 80)
    session_id = _clean_str(data.get('session_id'), 64)
    if token not in _screen_pending:
        return jsonify({"status": "error", "message": "QR الشاشة انتهت صلاحيته — اعمل refresh"}), 400
    if not _can_access_session(session['admin'], session_id):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403
    sess = sessions_col.find_one({"session_id": session_id, "is_open": True})
    if not sess:
        return jsonify({"status": "error", "message": "الجلسة غير موجودة أو مغلقة"}), 404
    secret = sess.get('screen_secret')
    if not secret:
        secret = generate_screen_secret()
        sessions_col.update_one({"session_id": session_id}, {"$set": {"screen_secret": secret}})
    _screen_pending[token]['session_id'] = session_id
    _screen_pending[token]['secret'] = secret
    _screen_pending[token]['session_info'] = {
        "subject_name": sess.get('subject_name', ''),
        "title": sess.get('title', ''),
        "type": sess.get('type', 'Lecture')
    }
    return jsonify({"status": "success", "message": "تم ربط الشاشة بالجلسة!"})


@app.route('/api/screen-pair', methods=['POST'])
@limiter.limit("20 per minute")
def screen_pair():
    data = request.get_json(force=True, silent=True) or {}
    code = _clean_str(data.get('screen_code'), 12)
    if len(code) != 8 or not code.isdigit():
        return jsonify({"status": "error", "message": "كود الشاشة يجب أن يكون 8 أرقام"}), 400
    open_sessions = list(sessions_col.find({"is_open": True}, {"_id": 0}))
    for sess in open_sessions:
        secret = sess.get('screen_secret')
        if not secret:
            continue
        if verify_screen_code(code, secret):
            return jsonify({"status": "success", "screen_secret": secret,
                "session": {"session_id": sess['session_id'],
                            "subject_name": sess.get('subject_name', ''),
                            "title": sess.get('title', ''),
                            "type": sess.get('type', 'Lecture'),
                            "created_by": sess.get('created_by', '')}})
    return jsonify({"status": "error", "message": "الكود انتهت صلاحيته أو غير صحيح"}), 404


# ============================================================
#  🖥️ screen-live-code (مع دعم الكويزات)
# ============================================================
@app.route('/api/screen-live-code')
@limiter.limit("300 per minute")
def screen_live_code():
    _auto_close_ta_sessions()
    secret = _clean_str(request.args.get('secret'), 80)
    legacy_code = _clean_str(request.args.get('screen_code'), 12)
    is_quiz = request.args.get('is_quiz') == '1'

    if is_quiz:
        if not secret:
            return jsonify({"status": "error", "message": "missing"}), 400

        quiz = quizzes_col.find_one({"id": secret})
        if not quiz:
            return jsonify({"status": "closed", "message": "الكويز غير موجود"})

        if not quiz.get('broadcast_active'):
            return jsonify({
                "status": "closed",
                "message": "البث مغلق",
                "session_info": {
                    "subject_name": quiz.get('subject_name', ''),
                    "title": quiz.get('title', ''),
                    "type": "Quiz",
                    "is_quiz": True
                }
            })

        now = time.time()
        m_step = int(now // MANUAL_STEP_INTERVAL)
        m_remaining = MANUAL_STEP_INTERVAL - (now - (m_step * MANUAL_STEP_INTERVAL))
        manual_code = get_quiz_manual_code(m_step, quiz.get('id', ''))

        q_step = int(now // QR_STEP_INTERVAL)
        q_remaining = QR_STEP_INTERVAL - (now - (q_step * QR_STEP_INTERVAL))
        qr_code = get_quiz_qr_code(q_step, quiz.get('id', ''))

        return jsonify({
            "status": "success",
            "session_info": {
                "subject_name": quiz.get('subject_name', ''),
                "title": quiz.get('title', ''),
                "type": "Quiz",
                "is_quiz": True
            },
            "manual_code": manual_code,
            "manual_remaining": m_remaining,
            "manual_interval": MANUAL_STEP_INTERVAL,
            "qr_code": qr_code,
            "qr_remaining": q_remaining,
            "qr_interval": QR_STEP_INTERVAL
        })

    sess = None
    if secret and len(secret) >= 16:
        sess = sessions_col.find_one({"screen_secret": secret}, {"_id": 0})
    elif legacy_code and len(legacy_code) == 8 and legacy_code.isdigit():
        open_sessions = list(sessions_col.find({"is_open": True}, {"_id": 0}))
        for s in open_sessions:
            sc = s.get('screen_secret')
            if sc and verify_screen_code(legacy_code, sc):
                sess = s
                break
        if not sess:
            return jsonify({"status": "closed", "message": "انتهت صلاحية الكود"})
    else:
        return jsonify({"status": "error", "message": "invalid"}), 400

    if not sess:
        return jsonify({"status": "closed", "message": "الجلسة انتهت"})
    if not sess.get('is_open'):
        return jsonify({"status": "closed", "message": "الجلسة مغلقة الآن",
            "session_info": {"subject_name": sess.get('subject_name', ''),
                             "title": sess.get('title', ''), "type": sess.get('type', 'Lecture')}})
    session_id = sess['session_id']
    now = time.time()
    m_step = int(now // MANUAL_STEP_INTERVAL)
    m_remaining = MANUAL_STEP_INTERVAL - (now - (m_step * MANUAL_STEP_INTERVAL))
    manual_code = get_step_code(m_step, session_id)
    q_step = int(now // QR_STEP_INTERVAL)
    q_remaining = QR_STEP_INTERVAL - (now - (q_step * QR_STEP_INTERVAL))
    qr_code = get_qr_step_code(q_step, session_id)
    return jsonify({"status": "success",
        "session_info": {"subject_name": sess.get('subject_name', ''),
                         "title": sess.get('title', ''), "type": sess.get('type', 'Lecture')},
        "manual_code": manual_code, "manual_remaining": m_remaining,
        "manual_interval": MANUAL_STEP_INTERVAL, "qr_code": qr_code,
        "qr_remaining": q_remaining, "qr_interval": QR_STEP_INTERVAL})


@app.route('/api/admin-action', methods=['POST'])
@limiter.limit("300 per hour")
def admin_action():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    curr_user = users_col.find_one({"username": session['admin']['username']})
    role = session['admin']['role']
    curr = session['admin']
    data = request.get_json(force=True, silent=True) or {}
    action = _clean_str(data.get('action'), 60)

    if action == 'change_my_password':
        new_pw = str(data.get('new_password', ''))
        if len(new_pw) < 8:
            return jsonify({"status": "error", "message": "كلمة المرور يجب أن تكون 8 أحرف على الأقل"}), 400
        users_col.update_one({"username": curr['username']}, {"$set": {"password": generate_password_hash(new_pw)}})
        return jsonify({"status": "success"})

    elif action == 'create_session':
        s_data = data.get('session', {}) or {}
        sub_id = _clean_str(s_data.get('subject_id'), 60)
        session_type = _clean_str(s_data.get('type', 'Lecture'), 20)
        custom_title = _clean_str(s_data.get('title', ''), 100)
        subject_name = _clean_str(s_data.get('subject_name', ''), 200)
        if session_type not in ('Lecture', 'Section'):
            return jsonify({"status": "error", "message": "نوع الجلسة غير صحيح"}), 400
        if not custom_title:
            return jsonify({"status": "error", "message": "يرجى إدخال عنوان الجلسة!"}), 400
        if role == 'super_admin':
            pass
        elif role == 'doctor':
            if sub_id not in curr_user.get('allowed_subjects', []):
                return jsonify({"status": "error", "message": "غير مصرح"}), 403
        elif role == 'ta':
            if session_type != 'Section':
                return jsonify({"status": "error", "message": "المعيد يمكنه فتح سكاشن فقط!"}), 403
            if sub_id not in curr_user.get('allowed_subjects', []):
                return jsonify({"status": "error", "message": "غير مصرح"}), 403
        else:
            return jsonify({"status": "error", "message": "غير مصرح"}), 403

        sess_id = f"SESS_{int(time.time() * 1000)}"
        new_session = {"session_id": sess_id, "subject_id": sub_id,
                       "subject_name": subject_name, "type": session_type,
                       "title": custom_title, "is_open": True,
                       "created_by": curr['name'], "created_by_username": curr['username'],
                       "creator_role": role, "created_at": _cairo_now(),
                       "created_at_dt": datetime.utcnow(),
                       "screen_secret": generate_screen_secret()}
        sessions_col.insert_one(new_session)
        safe_session = {k: v for k, v in new_session.items() if k not in ('screen_secret', '_id', 'created_at_dt')}
        return jsonify({"status": "success", "session": safe_session})

    elif action == 'update_session_title':
        sess_id = _clean_str(data.get('session_id'), 64)
        new_title = _clean_str(data.get('new_title'), 100)
        if not sess_id or not new_title:
            return jsonify({"status": "error", "message": "بيانات ناقصة!"}), 400
        if not _can_access_session(curr, sess_id):
            return jsonify({"status": "error", "message": "غير مصرح!"}), 403
        sessions_col.update_one({"session_id": sess_id}, {"$set": {"title": new_title}})
        return jsonify({"status": "success", "message": "تم تحديث العنوان"})

    elif action == 'toggle_session':
        sess_id = _clean_str(data.get('session_id'), 64)
        if not _can_access_session(curr, sess_id):
            return jsonify({"status": "error", "message": "غير مصرح!"}), 403
        is_open = bool(data.get('is_open'))
        if is_open:
            sessions_col.update_one({"session_id": sess_id},
                {"$set": {"is_open": True, "created_at_dt": datetime.utcnow(),
                          "reopened_at": _cairo_now()},
                 "$unset": {"auto_closed": "", "closed_at": ""}})
        else:
            sessions_col.update_one({"session_id": sess_id},
                {"$set": {"is_open": False, "closed_at": _cairo_now(),
                          "closed_by": curr.get('username', 'manual')}})
        return jsonify({"status": "success"})

    elif action == 'delete_session':
        sess_id = _clean_str(data.get('session_id'), 64)
        if not _can_access_session(curr, sess_id):
            return jsonify({"status": "error", "message": "غير مصرح!"}), 403
        sessions_col.delete_one({"session_id": sess_id})
        attendance_col.delete_many({"session_id": sess_id})
        return jsonify({"status": "success"})

    elif action == 'add_manual_attendance':
        sess_id = _clean_str(data.get('session_id'), 64)
        s_id = _clean_str(data.get('student_id'), 20)
        s_name = _clean_str(data.get('student_name'), 100)
        if len(s_id) != 7 or not s_id.isdigit():
            return jsonify({"status": "error", "message": "كود الطالب غير صحيح!"}), 400
        if not _can_access_session(curr, sess_id):
            return jsonify({"status": "error", "message": "غير مصرح!"}), 403
        sess = sessions_col.find_one({"session_id": sess_id})
        if not sess:
            return jsonify({"status": "error", "message": "الجلسة غير موجودة!"}), 404
        if attendance_col.find_one({"student_id": s_id, "session_id": sess_id}):
            return jsonify({"status": "error", "message": "الطالب مسجل حضوره بالفعل!"}), 409
        st_record = students_col.find_one({"student_id": s_id})
        year = st_record.get('year', '') if st_record else ''
        dept = st_record.get('department', '') if st_record else ''
        record = {"student_id": s_id, "student_name": s_name, "year": year, "department": dept,
                  "subject_id": sess['subject_id'], "subject_name": sess['subject_name'],
                  "session_id": sess_id, "session_title": sess['title'],
                  "session_type": sess.get('type', 'Lecture'), "timestamp": _cairo_now(),
                  "device_token": "MANUAL_BY_ADMIN", "ip": "ADMIN",
                  "is_manual": True, "scan_type": "admin"}
        attendance_col.insert_one(record)
        record.pop('_id', None)
        return jsonify({"status": "success", "message": "تم التحضير!", "record": record})

    elif action == 'edit_attendance_record':
        sess_id = _clean_str(data.get('session_id'), 64)
        old_id = _clean_str(data.get('old_id'), 20)
        new_id = _clean_str(data.get('new_id'), 20)
        new_name = _clean_str(data.get('new_name'), 100)
        if len(new_id) != 7 or not new_id.isdigit():
            return jsonify({"status": "error", "message": "الكود الجديد غير صحيح"}), 400
        attendance_col.update_one({"session_id": sess_id, "student_id": old_id},
            {"$set": {"student_id": new_id, "student_name": new_name}})
        return jsonify({"status": "success"})

    elif action == 'delete_attendance_record':
        attendance_col.delete_one({"student_id": _clean_str(data.get('student_id'), 20),
                                    "session_id": _clean_str(data.get('session_id'), 64)})
        return jsonify({"status": "success"})

    elif action == 'manage_subject' and role == 'super_admin':
        sub_act = _clean_str(data.get('sub'), 20)
        if sub_act == 'add':
            s_obj = data.get('subject', {}) or {}
            sub_id = f"SUB_{int(time.time())}"
            subjects_col.insert_one({"id": sub_id, "name": _clean_str(s_obj.get('name', ''), 200),
                "year": _clean_str(s_obj.get('year', ''), 40),
                "department": _clean_str(s_obj.get('department', ''), 60),
                "image": str(s_obj.get('image', ''))[:500000], "added_by": curr['name']})
        elif sub_act == 'edit':
            sub_id = _clean_str(data.get('id', ''), 60)
            s_obj = data.get('subject', {}) or {}
            target = subjects_col.find_one({"id": sub_id})
            if not target:
                return jsonify({"status": "error", "message": "المادة غير موجودة!"}), 404
            new_name = _clean_str(s_obj.get('name', target.get('name', '')), 200)
            updates = {"name": new_name,
                       "year": _clean_str(s_obj.get('year', target.get('year', '')), 40),
                       "department": _clean_str(s_obj.get('department', target.get('department', '')), 60)}
            if s_obj.get('image'):
                updates['image'] = str(s_obj.get('image'))[:500000]
            subjects_col.update_one({"id": sub_id}, {"$set": updates})
            if new_name and new_name != target.get('name'):
                sessions_col.update_many({"subject_id": sub_id}, {"$set": {"subject_name": new_name}})
                attendance_col.update_many({"subject_id": sub_id}, {"$set": {"subject_name": new_name}})
        elif sub_act == 'delete':
            sub_id = _clean_str(data.get('id', ''), 60)
            subjects_col.delete_one({"id": sub_id})
            sessions_col.delete_many({"subject_id": sub_id})
            attendance_col.delete_many({"subject_id": sub_id})
        return jsonify({"status": "success"})

    elif action == 'manage_staff':
        if role == 'ta':
            return jsonify({"status": "error", "message": "المعيد ليس له صلاحية!"}), 403
        sub_act = _clean_str(data.get('sub', ''), 20)
        staff_data = data.get('staff', {}) or {}
        if sub_act == 'add':
            new_role = _clean_str(staff_data.get('role', 'ta'), 20)
            if new_role not in ('ta', 'doctor'):
                new_role = 'ta'
            allowed_subjects = staff_data.get('allowed_subjects', [])
            if not isinstance(allowed_subjects, list):
                allowed_subjects = []
            allowed_subjects = [_clean_str(s, 60) for s in allowed_subjects][:200]
            if role == 'doctor':
                new_role = 'ta'
                my_subs = curr_user.get('allowed_subjects', [])
                if not all(s in my_subs for s in allowed_subjects):
                    return jsonify({"status": "error", "message": "غير مصرح"}), 403
            target_user = _clean_str(staff_data.get('username', ''), 100)
            password = str(staff_data.get('password', ''))
            name = _clean_str(staff_data.get('name', ''), 100)
            if not target_user or not password or not name:
                return jsonify({"status": "error", "message": "بيانات ناقصة"}), 400
            if len(password) < 6:
                return jsonify({"status": "error", "message": "كلمة المرور 6 أحرف على الأقل"}), 400
            if users_col.find_one({"username": target_user}):
                return jsonify({"status": "error", "message": "اسم المستخدم موجود مسبقاً!"}), 400
            users_col.insert_one({"name": name, "username": target_user,
                "password": generate_password_hash(password),
                "role": new_role, "allowed_subjects": allowed_subjects,
                "is_active": True, "created_by": curr['username'],
                "created_at": datetime.utcnow()})
        elif sub_act == 'delete':
            target_username = _clean_str(data.get('username', ''), 100)
            if target_username == curr['username']:
                return jsonify({"status": "error", "message": "لا يمكنك حذف حسابك!"}), 403
            target = users_col.find_one({"username": target_username})
            if not target:
                return jsonify({"status": "error", "message": "العضو غير موجود!"}), 404
            if role == 'doctor':
                if target.get('role') != 'ta':
                    return jsonify({"status": "error", "message": "يمكنك حذف المعيدين فقط!"}), 403
                my_subs = curr_user.get('allowed_subjects', [])
                target_subs = target.get('allowed_subjects', [])
                has_common = any(s in my_subs for s in target_subs)
                created_by_him = (target.get('created_by') == curr['username'])
                if not (has_common or created_by_him):
                    return jsonify({"status": "error", "message": "غير مصرح!"}), 403
            users_col.delete_one({"username": target_username})
        return jsonify({"status": "success"})

    elif action == 'edit_staff':
        target_username = _clean_str(data.get('target_username', ''), 100)
        if not target_username:
            return jsonify({"status": "error", "message": "بيانات ناقصة"}), 400
        target = users_col.find_one({"username": target_username})
        if not target:
            return jsonify({"status": "error", "message": "العضو غير موجود!"}), 404
        if target_username == curr['username']:
            return jsonify({"status": "error", "message": "لا يمكنك تعديل حسابك!"}), 403
        staff_data = data.get('staff', {}) or {}
        updates = {}
        if role == 'super_admin':
            new_name = _clean_str(staff_data.get('name', target.get('name', '')), 100)
            new_role = _clean_str(staff_data.get('role', target.get('role', 'ta')), 20)
            new_subs = staff_data.get('allowed_subjects', target.get('allowed_subjects', []))
            if not isinstance(new_subs, list):
                new_subs = []
            new_subs = [_clean_str(s, 60) for s in new_subs][:200]
            if new_name:
                updates['name'] = new_name
            if new_role in ('doctor', 'ta'):
                updates['role'] = new_role
            updates['allowed_subjects'] = new_subs
        elif role == 'doctor':
            if target.get('role') != 'ta':
                return jsonify({"status": "error", "message": "يمكنك تعديل المعيدين فقط!"}), 403
            my_subs = curr_user.get('allowed_subjects', [])
            target_subs = target.get('allowed_subjects', [])
            has_common = any(s in my_subs for s in target_subs)
            created_by_him = (target.get('created_by') == curr['username'])
            if not (has_common or created_by_him):
                return jsonify({"status": "error", "message": "غير مصرح!"}), 403
            new_subs = staff_data.get('allowed_subjects', [])
            if not isinstance(new_subs, list):
                new_subs = []
            new_subs = [_clean_str(s, 60) for s in new_subs][:200]
            if not all(s in my_subs for s in new_subs):
                return jsonify({"status": "error", "message": "غير مصرح"}), 403
            other_subs = [s for s in target_subs if s not in my_subs]
            updates['allowed_subjects'] = list(set(other_subs + new_subs))
        else:
            return jsonify({"status": "error", "message": "غير مصرح"}), 403
        if updates:
            users_col.update_one({"username": target_username}, {"$set": updates})
        return jsonify({"status": "success", "message": "تم التحديث"})

    elif action == 'toggle_staff_status':
        target_username = _clean_str(data.get('target_username', ''), 100)
        if not target_username:
            return jsonify({"status": "error", "message": "بيانات ناقصة"}), 400
        target = users_col.find_one({"username": target_username})
        if not target:
            return jsonify({"status": "error", "message": "العضو غير موجود!"}), 404
        if target_username == curr['username']:
            return jsonify({"status": "error", "message": "لا يمكنك إيقاف حسابك!"}), 403
        if role == 'doctor':
            if target.get('role') != 'ta':
                return jsonify({"status": "error", "message": "يمكنك التحكم في المعيدين فقط!"}), 403
            my_subs = curr_user.get('allowed_subjects', [])
            target_subs = target.get('allowed_subjects', [])
            has_common = any(s in my_subs for s in target_subs)
            created_by_him = (target.get('created_by') == curr['username'])
            if not (has_common or created_by_him):
                return jsonify({"status": "error", "message": "غير مصرح!"}), 403
        elif role != 'super_admin':
            return jsonify({"status": "error", "message": "غير مصرح"}), 403
        new_status = bool(data.get('is_active', True))
        users_col.update_one({"username": target_username}, {"$set": {"is_active": new_status}})
        return jsonify({"status": "success", "message": "تم التنشيط" if new_status else "تم الإيقاف"})

    elif action == 'wipe_all' and role == 'super_admin':
        provided_pw = str(data.get('admin_password', ''))
        admin_doc = users_col.find_one({"username": ADMIN_USERNAME})
        ok = False
        if admin_doc:
            try:
                ok = check_password_hash(admin_doc.get('password', ''), provided_pw)
            except Exception:
                ok = False
        if ok:
            attendance_col.delete_many({})
            sessions_col.delete_many({})
            return jsonify({"status": "success", "message": "تم التصفير بنجاح"})
        return jsonify({"status": "error", "message": "كلمة المرور غير صحيحة!"}), 403

    return jsonify({"status": "success"})


# ============================================================
#  📝 QUIZ SYSTEM
# ============================================================
def _quiz_can_manage(admin: dict, quiz: dict) -> bool:
    if not quiz:
        return False
    role = admin.get('role')
    if role == 'super_admin':
        return True
    curr = users_col.find_one({"username": admin.get('username')})
    if not curr:
        return False
    if quiz.get('subject_id') not in (curr.get('allowed_subjects') or []):
        return False
    if role == 'doctor':
        return True
    if role == 'ta':
        return quiz.get('created_by_username') == curr.get('username')
    return False


def _quiz_is_teacher_of_subject(admin: dict, subject_id: str) -> bool:
    if admin.get('role') == 'super_admin':
        return True
    curr = users_col.find_one({"username": admin.get('username')})
    return subject_id in (curr.get('allowed_subjects') or []) if curr else False


def _normalize_word_text(text):
    text = text.replace('\ufeff', '').replace('\xa0', ' ')
    text = re.sub(r'([^\s\n])([A-D][\.:])', r'\1\n\2', text)
    text = re.sub(r'(\s)([A-D][\.:])', r'\1\n\2', text)
    text = re.sub(r'([^\s\n])\s*(Correct Answer:)', r'\1\n\2', text, flags=re.IGNORECASE)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)


def _extract_docx_text(file_bytes):
    text_result = ""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        text_result = "\n".join(parts)
    except Exception:
        pass
    if not text_result:
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                xml_content = z.read('word/document.xml')
            root = ET.fromstring(xml_content)
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            parts = []
            for p in root.findall('.//w:p', ns):
                p_text = "".join([t.text for t in p.findall('.//w:t', ns) if t.text])
                if p_text:
                    parts.append(p_text)
            text_result = "\n".join(parts)
        except Exception:
            pass
    return text_result


def _parse_word_questions(text):
    if not text:
        return []
    text = _normalize_word_text(text)
    lines = text.splitlines()
    questions = []
    current_q = None
    q_pattern = re.compile(r'^(\d+)[\.:\)]\s*(.+)')
    for line in lines:
        line = line.strip()
        if not line:
            continue
        q_match = q_pattern.match(line)
        if q_match:
            if current_q:
                questions.append(current_q)
            current_q = {'text': q_match.group(2), 'options': [], 'correct': 'a', 'type': 'mcq'}
            continue
        if current_q:
            opt_match = re.match(r'^([A-D])[\.:\)]\s*(.+)', line, re.IGNORECASE)
            if opt_match:
                current_q['options'].append(opt_match.group(2))
                continue
            ans_match = re.match(r'^Correct Answer:\s*([A-D])', line, re.IGNORECASE)
            if ans_match:
                current_q['correct'] = ans_match.group(1).lower()
                continue
            if line.lower() in ['true', 'false']:
                current_q['type'] = 'tf'
                current_q['options'] = ["True", "False", "", ""]
                current_q['correct'] = 'a' if line.lower() == 'true' else 'b'
                continue
    if current_q:
        questions.append(current_q)
    cleaned = []
    for q in questions:
        if q['type'] == 'tf':
            q['options'] = ["True", "False", "", ""]
        else:
            opts = [o for o in q['options'] if o]
            if len(opts) < 2:
                continue
            while len(q['options']) < 4:
                q['options'].append('')
            q['options'] = q['options'][:4]
        cleaned.append(q)
    return cleaned


def _generate_attempt_questions(quiz):
    generated = []
    sections = quiz.get('sections_config')
    if sections:
        counter = 1
        for section in sections:
            all_qs = section.get('all_questions', [])
            num = int(section.get('num_to_pick') or 0)
            mcq_pts = float(section.get('mcq_points') or 1)
            tf_pts = float(section.get('tf_points') or 1)
            if not all_qs or num < 1:
                continue
            actual_num = min(num, len(all_qs))
            picked = random.sample(all_qs, actual_num)
            for q in picked:
                pts = tf_pts if q.get('type') == 'tf' else mcq_pts
                generated.append({
                    "id": f"q{counter}", "text": q.get('text', ''),
                    "type": q.get('type', 'mcq'), "options": q.get('options', []),
                    "correct": q.get('correct', 'a'), "points": pts,
                    "section": section.get('name', '')
                })
                counter += 1
        random.shuffle(generated)
    else:
        generated = list(quiz.get('questions', []))
        random.shuffle(generated)
    return generated


def _compute_expected_count(quiz):
    if quiz.get('sections_config'):
        return sum(min(s.get('num_to_pick', 0), len(s.get('all_questions', [])))
                   for s in quiz['sections_config'])
    return len(quiz.get('questions', []))


def _student_can_access_quiz(student: dict, quiz: dict) -> bool:
    if not student or not quiz:
        return False
    subject = subjects_col.find_one({"id": quiz.get('subject_id')})
    if not subject:
        return False
    s_year = student.get('year') or ''
    s_dept = student.get('department') or 'عام (IT)'
    subj_year = subject.get('year') or ''
    subj_dept = subject.get('department') or 'عام (IT)'
    if subj_year != s_year:
        return False
    if subj_dept not in (s_dept, "عام (IT)"):
        return False
    return True


# ============================================================
#  🎯 AUTO-GRADE HELPER (used by time expiry + violation fallback)
# ============================================================
def _auto_grade_attempt(attempt, quiz):
    """
    Grades an in-progress attempt using whatever answers are saved.
    Used when:
      - Server detects expires_at has passed
      - A "violation" arrives after expiry
      - Submit arrives late
    Sets status='completed' with real score (NOT zero, NOT cheat).
    """
    attempt_questions = attempt.get('generated_questions') or quiz.get('questions', [])
    answers = attempt.get('answers') or {}
    score = 0.0
    total = 0.0
    details = []
    for q in attempt_questions:
        total += float(q.get('points', 1))
        ans = answers.get(q['id'])
        ok = (ans == q['correct'])
        if ok:
            score += float(q.get('points', 1))
        details.append({"qid": q['id'], "text": q['text'], "user": ans,
                        "correct": q['correct'], "ok": ok, "points": q.get('points', 1)})
    quiz_attempts_col.update_one(
        {"attempt_id": attempt['attempt_id'], "status": "in_progress"},
        {"$set": {
            "status": "completed",
            "score": round(score, 2),
            "total_points": round(total, 2),
            "details": details,
            "auto_submitted": True,
            "auto_reason": "time_expired",
            "submitted_at": _cairo_now(),
            "submitted_at_dt": datetime.utcnow()
        }})
    return score, total


# ============================================================
#  🎫 QUIZ LIVE CODES
# ============================================================
def get_quiz_manual_code(step, quiz_id):
    return _generate_code(QR_SECRET_SALT, step, f"quiz_m_{quiz_id}")


def get_quiz_qr_code(step, quiz_id):
    return _generate_code(QR_SECRET_SALT, step, f"quiz_q_{quiz_id}")


def verify_quiz_code(user_code, quiz_id, scan_type='manual'):
    clean = (user_code or '').strip().upper()
    if not clean or len(clean) != 6:
        return False
    now = time.time()
    if scan_type == 'qr':
        current_step = int(now // QR_STEP_INTERVAL)
        return _safe_eq(clean, get_quiz_qr_code(current_step, quiz_id))
    current_step = int(now // MANUAL_STEP_INTERVAL)
    if _safe_eq(clean, get_quiz_manual_code(current_step, quiz_id)):
        return True
    elapsed = now - (current_step * MANUAL_STEP_INTERVAL)
    if elapsed < MANUAL_GRACE_PERIOD and _safe_eq(clean, get_quiz_manual_code(current_step - 1, quiz_id)):
        return True
    return False


@app.route('/api/quiz-live-code')
def api_quiz_live_code():
    quiz_id = _clean_str(request.args.get('quiz_id'), 60)
    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz:
        return jsonify({"status": "error", "message": "not found"}), 404
    if not quiz.get('broadcast_active'):
        return jsonify({"status": "closed", "message": "البث مغلق"})

    now = time.time()
    m_step = int(now // MANUAL_STEP_INTERVAL)
    m_remaining = MANUAL_STEP_INTERVAL - (now - (m_step * MANUAL_STEP_INTERVAL))
    manual_code = get_quiz_manual_code(m_step, quiz_id)

    q_step = int(now // QR_STEP_INTERVAL)
    q_remaining = QR_STEP_INTERVAL - (now - (q_step * QR_STEP_INTERVAL))
    qr_code = get_quiz_qr_code(q_step, quiz_id)

    completed = quiz_attempts_col.count_documents({"quiz_id": quiz_id, "status": "completed"})
    in_progress = quiz_attempts_col.count_documents({"quiz_id": quiz_id, "status": "in_progress"})

    return jsonify({
        "status": "success",
        "quiz": {
            "id": quiz['id'], "title": quiz.get('title', ''),
            "subject_name": quiz.get('subject_name', ''),
            "duration": quiz.get('duration', 0),
            "expected_questions": _compute_expected_count(quiz)
        },
        "manual_code": manual_code, "manual_remaining": m_remaining,
        "manual_interval": MANUAL_STEP_INTERVAL,
        "qr_code": qr_code, "qr_remaining": q_remaining,
        "qr_interval": QR_STEP_INTERVAL,
        "attempt_count": completed + in_progress,
        "completed_count": completed
    })


@app.route('/api/admin-toggle-quiz-broadcast', methods=['POST'])
def api_admin_toggle_quiz_broadcast():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    quiz_id = _clean_str(data.get('id'), 60)
    active = bool(data.get('active'))
    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403
    updates = {"broadcast_active": active}
    if active:
        updates['broadcast_started_at'] = _cairo_now()
    quizzes_col.update_one({"id": quiz_id}, {"$set": updates})
    return jsonify({"status": "success", "active": active})


# ============================================================
#  🚀 QUIZ QUICK START (with strict checks + auto-unblock)
# ============================================================
@app.route('/api/quiz-quick-start', methods=['POST'])
def api_quiz_quick_start():
    data = request.get_json(force=True, silent=True) or {}
    quiz_id = _clean_str(data.get('quiz_id'), 60)
    student_id = _clean_str(data.get('student_id'), 20)
    student_name = _clean_str(data.get('student_name'), 100)
    code = _clean_str(data.get('code'), 12).upper()
    scan_type = _clean_str(data.get('scan_type', 'manual'), 20).lower()

    if len(student_id) != 7 or not student_id.isdigit():
        return jsonify({"status": "error", "message": "ID غير صحيح"}), 400

    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz:
        return jsonify({"status": "error", "message": "الكويز غير موجود"}), 404

    student = students_col.find_one({"student_id": student_id})
    if not student:
        return jsonify({"status": "error", "message": "الطالب غير مسجل"}), 404

    if not _student_can_access_quiz(student, quiz):
        return jsonify({"status": "error", "message": "هذا الامتحان غير متاح لمستواك الدراسي"}), 403

    if not quiz.get('is_open'):
        return jsonify({"status": "error", "message": "الكويز مغلق"}), 400
    if not quiz.get('broadcast_active'):
        return jsonify({"status": "error", "message": "الامتحان غير مفتوح حالياً"}), 400

    if not verify_quiz_code(code, quiz_id, scan_type):
        return jsonify({"status": "error", "message": "الكود غير صحيح أو انتهت صلاحيته"}), 400

    # ⚡ Auto-unblock for transient disconnect
    block = quiz_blocks_col.find_one({"quiz_id": quiz_id, "student_id": student_id, "unblocked_at": None})
    if block:
        reason = (block.get('reason') or '').strip()
        blocked_dt = block.get('blocked_at_dt')
        if reason in ('انقطاع', 'Session ended from server', 'انقطاع شبكة') and blocked_dt:
            try:
                elapsed = (datetime.utcnow() - blocked_dt).total_seconds()
                if elapsed < 60:
                    quiz_blocks_col.delete_many({"quiz_id": quiz_id, "student_id": student_id})
                    quiz_attempts_col.delete_many({
                        "quiz_id": quiz_id, "student_id": student_id,
                        "status": "cancelled"
                    })
                    block = None
            except Exception:
                pass
        if block:
            return jsonify({"status": "blocked", "message": "محظور"}), 403

    existing = quiz_attempts_col.find_one({"quiz_id": quiz_id, "student_id": student_id})
    if existing:
        if existing.get('status') == 'completed':
            return jsonify({"status": "already_done",
                            "score": existing.get('score'),
                            "total": existing.get('total_points')})
        if existing.get('status') == 'in_progress':
            last_hb = existing.get('last_heartbeat_dt') or existing.get('started_at')
            if last_hb and (datetime.utcnow() - last_hb).total_seconds() > HEARTBEAT_GRACE_SEC:
                quiz_attempts_col.update_one({"attempt_id": existing['attempt_id']},
                    {"$set": {"status": "cancelled", "cancelled_reason": "انقطاع",
                              "cancelled_at": _cairo_now()}})
                quiz_blocks_col.update_one({"quiz_id": quiz_id, "student_id": student_id},
                    {"$setOnInsert": {"quiz_id": quiz_id, "student_id": student_id,
                                      "student_name": student_name, "reason": "انقطاع",
                                      "blocked_at": _cairo_now(),
                                      "blocked_at_dt": datetime.utcnow(),
                                      "unblocked_at": None}}, upsert=True)
                return jsonify({"status": "blocked", "message": "تم الحظر"}), 403
            if existing.get('generated_questions'):
                generated_qs = existing['generated_questions']
                total_pts = existing.get('total_points', 0)
            else:
                generated_qs = _generate_attempt_questions(quiz)
                total_pts = round(sum(q.get('points', 1) for q in generated_qs), 2)
                quiz_attempts_col.update_one({"attempt_id": existing['attempt_id']},
                    {"$set": {"generated_questions": generated_qs, "total_points": total_pts}})
            attempt = existing
        else:
            # Cancelled but not yet reblocked — reset
            quiz_attempts_col.delete_one({"attempt_id": existing['attempt_id']})
            return api_quiz_quick_start()
    else:
        generated_qs = _generate_attempt_questions(quiz)
        total_pts = round(sum(q.get('points', 1) for q in generated_qs), 2)
        attempt_id = f"AT_{int(time.time()*1000)}_{secrets.token_hex(3)}"
        now = datetime.utcnow()
        quiz_attempts_col.insert_one({
            "attempt_id": attempt_id, "quiz_id": quiz_id, "student_id": student_id,
            "student_name": student_name, "started_at": now, "last_heartbeat_dt": now,
            "expires_at": now + timedelta(minutes=quiz.get('duration', 30)),
            "status": "in_progress", "answers": {}, "score": None,
            "generated_questions": generated_qs,
            "total_points": total_pts, "violations": 0
        })
        attempt = quiz_attempts_col.find_one({"attempt_id": attempt_id})

    public_qs = [{"id": q['id'], "text": q['text'], "type": q['type'],
                  "options": q['options'], "points": q['points']} for q in generated_qs]
    expires_in = max(0, int((attempt['expires_at'] - datetime.utcnow()).total_seconds()))

    return jsonify({
        "status": "success",
        "attempt_id": attempt['attempt_id'],
        "quiz": {
            "id": quiz['id'], "title": quiz['title'],
            "duration": quiz['duration'], "questions": public_qs,
            "total_points": total_pts
        },
        "expires_in": expires_in,
        "server_now": datetime.utcnow().isoformat() + "Z"
    })


@app.route('/api/quiz-screen-pair-token', methods=['POST'])
@limiter.limit("30 per minute")
def api_quiz_screen_pair_token():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401

    data = request.get_json(force=True, silent=True) or {}
    token = _clean_str(data.get('token'), 80)
    quiz_id = _clean_str(data.get('quiz_id'), 60)

    if token not in _screen_pending:
        return jsonify({"status": "error", "message": "QR الشاشة انتهت صلاحيته — اعمل refresh للشاشة"}), 400

    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403

    if not quiz.get('broadcast_active'):
        return jsonify({"status": "error", "message": "يجب بدء البث أولاً"}), 400

    _screen_pending[token]['session_id'] = quiz_id
    _screen_pending[token]['secret'] = quiz_id
    _screen_pending[token]['session_info'] = {
        "subject_name": quiz.get('subject_name', ''),
        "title": quiz.get('title', ''),
        "type": "Quiz",
        "is_quiz": True,
        "quiz_id": quiz_id
    }

    return jsonify({"status": "success", "message": "تم ربط الشاشة بالامتحان!"})


@app.route('/api/quiz-by-subject')
def api_quiz_by_subject():
    sid = _clean_str(request.args.get('student_id'), 20)
    subject_id = _clean_str(request.args.get('subject_id'), 60)
    if len(sid) != 7 or not sid.isdigit():
        return jsonify({"status": "error", "message": "ID غير صحيح"}), 400

    student = students_col.find_one({"student_id": sid})
    if not student:
        return jsonify({"status": "success", "quizzes": []})

    subject = subjects_col.find_one({"id": subject_id})
    if not subject:
        return jsonify({"status": "error", "message": "المادة غير موجودة"}), 404

    s_year = student.get('year') or ''
    s_dept = student.get('department') or 'عام (IT)'
    subj_year = subject.get('year') or ''
    subj_dept = subject.get('department') or 'عام (IT)'

    if subj_year != s_year:
        return jsonify({"status": "success", "quizzes": []})
    if subj_dept not in (s_dept, "عام (IT)"):
        return jsonify({"status": "success", "quizzes": []})

    items = list(quizzes_col.find(
        {"subject_id": subject_id, "is_open": True},
        {"questions.correct": 0, "_id": 0, "sections_config": 0}
    ).sort("created_at_dt", -1))

    for it in items:
        att = quiz_attempts_col.find_one({"quiz_id": it['id'], "student_id": sid}, {"_id": 0})
        block = quiz_blocks_col.find_one({"quiz_id": it['id'], "student_id": sid, "unblocked_at": None})
        it['my_status'] = 'blocked' if block else (att.get('status', 'not_started') if att else 'not_started')
        it['questions_count'] = _compute_expected_count(it)
        it['broadcast_active'] = it.get('broadcast_active', False)
        it.pop('questions', None)
        if att and att.get('status') == 'completed':
            it['my_score'] = att.get('score')
            it['my_total'] = att.get('total_points')

    return jsonify({"status": "success", "quizzes": items})


@app.route('/exam')
def exam_page_route():
    return render_template('exam.html')


@app.route('/api/admin-quizzes')
def api_admin_quizzes():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    role = session['admin']['role']
    curr = users_col.find_one({"username": session['admin']['username']},
                              {"allowed_subjects": 1})
    allowed = (curr.get('allowed_subjects') or []) if curr else []

    q = {}
    if role != 'super_admin':
        q = {"subject_id": {"$in": allowed}}
        if role == 'ta':
            q["created_by_username"] = session['admin']['username']

    items = list(quizzes_col.find(q, {
        "_id": 0,
        "sections_config.all_questions": 0
    }).sort("created_at_dt", -1))

    if not items:
        return jsonify({"status": "success", "quizzes": []})

    quiz_ids = [it['id'] for it in items]

    attempts_map = {}
    try:
        for row in quiz_attempts_col.aggregate([
            {"$match": {"quiz_id": {"$in": quiz_ids}, "status": "completed"}},
            {"$group": {"_id": "$quiz_id", "count": {"$sum": 1}}}
        ]):
            attempts_map[row['_id']] = row['count']
    except Exception:
        pass

    blocks_map = {}
    try:
        for row in quiz_blocks_col.aggregate([
            {"$match": {"quiz_id": {"$in": quiz_ids}, "unblocked_at": None}},
            {"$group": {"_id": "$quiz_id", "count": {"$sum": 1}}}
        ]):
            blocks_map[row['_id']] = row['count']
    except Exception:
        pass

    for it in items:
        it['expected_questions'] = (
            it.get('expected_questions') or _compute_expected_count(it)
        )
        it['attempts_count'] = attempts_map.get(it['id'], 0)
        it['blocked_count'] = blocks_map.get(it['id'], 0)
        it['broadcast_active'] = it.get('broadcast_active', False)

    return jsonify({"status": "success", "quizzes": items})


@app.route('/api/admin-quiz-details')
def api_admin_quiz_details():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    qid = _clean_str(request.args.get('id'), 60)
    quiz = quizzes_col.find_one({"id": qid}, {"_id": 0})
    if not quiz:
        return jsonify({"status": "error", "message": "غير موجود"}), 404
    if not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403

    if quiz.get('sections_config'):
        flat = []
        for si, section in enumerate(quiz['sections_config']):
            for qi, q in enumerate(section.get('all_questions', [])):
                pts = section['tf_points'] if q.get('type') == 'tf' else section['mcq_points']
                flat.append({"id": f"s{si}_q{qi}", "text": q['text'], "type": q['type'],
                             "options": q['options'], "correct": q['correct'],
                             "points": pts, "section": section.get('name', '')})
        quiz['_admin_flat_questions'] = flat
        quiz['_sections_info'] = [{"name": s.get('name', ''), "pool_size": len(s.get('all_questions', [])),
                                    "pick_count": s.get('num_to_pick', 0)}
                                   for s in quiz['sections_config']]
        quiz['_is_random_mode'] = True
        quiz['_expected_questions'] = _compute_expected_count(quiz)
    else:
        quiz['_is_random_mode'] = False
        quiz['_expected_questions'] = len(quiz.get('questions', []))

    quiz['broadcast_active'] = quiz.get('broadcast_active', False)

    attempts = list(quiz_attempts_col.find({"quiz_id": qid}, {"_id": 0}).sort("started_at", -1).limit(500))
    blocks = list(quiz_blocks_col.find({"quiz_id": qid, "unblocked_at": None}, {"_id": 0}).sort("blocked_at", -1))
    return jsonify({"status": "success", "quiz": quiz, "attempts": attempts, "blocks": blocks})


@app.route('/api/admin-delete-quiz', methods=['POST'])
def api_admin_delete_quiz():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    qid = _clean_str(data.get('id'), 60)
    quiz = quizzes_col.find_one({"id": qid})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403
    quizzes_col.delete_one({"id": qid})
    quiz_attempts_col.delete_many({"quiz_id": qid})
    quiz_blocks_col.delete_many({"quiz_id": qid})
    return jsonify({"status": "success"})


@app.route('/api/admin-toggle-quiz', methods=['POST'])
def api_admin_toggle_quiz():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    qid = _clean_str(data.get('id'), 60)
    is_open = bool(data.get('is_open'))
    quiz = quizzes_col.find_one({"id": qid})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403
    quizzes_col.update_one({"id": qid}, {"$set": {"is_open": is_open}})
    return jsonify({"status": "success"})


@app.route('/api/admin-quiz-unblock', methods=['POST'])
def api_admin_quiz_unblock():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    qid = _clean_str(data.get('quiz_id'), 60)
    sid = _clean_str(data.get('student_id'), 20)
    quiz = quizzes_col.find_one({"id": qid})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403
    quiz_blocks_col.delete_many({"quiz_id": qid, "student_id": sid})
    quiz_attempts_col.delete_many({"quiz_id": qid, "student_id": sid})
    return jsonify({"status": "success"})


# ============================================================
#  🔄 QUIZ RETAKE — full reset for a single student
# ============================================================
@app.route('/api/admin-quiz-retake', methods=['POST'])
def api_admin_quiz_retake():
    """
    إعادة الامتحان لطالب:
    - يحذف محاولته السابقة
    - يفك الحظر لو موجود
    - الطالب يرجع يقدر يدخل من جديد
    """
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401

    data = request.get_json(force=True, silent=True) or {}
    qid = _clean_str(data.get('quiz_id'), 60)
    sid = _clean_str(data.get('student_id'), 20)

    if not qid or not sid:
        return jsonify({"status": "error", "message": "بيانات ناقصة"}), 400

    quiz = quizzes_col.find_one({"id": qid})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403

    # ✅ Delete attempt + block (student can start fresh)
    quiz_attempts_col.delete_many({"quiz_id": qid, "student_id": sid})
    quiz_blocks_col.delete_many({"quiz_id": qid, "student_id": sid})

    return jsonify({
        "status": "success",
        "message": "تم إعادة الامتحان — الطالب يقدر يدخل من جديد"
    })


@app.route('/api/admin-parse-word', methods=['POST'])
def api_admin_parse_word():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    files_data = data.get('files', [])
    if not files_data or not isinstance(files_data, list):
        return jsonify({"status": "error", "message": "ارفع ملف Word واحد على الأقل"}), 400
    results = {}
    total_found = 0
    for file_entry in files_data:
        filename = _clean_str(file_entry.get('name', 'unknown'), 200)
        content = file_entry.get('content', '')
        try:
            if ',' in content:
                content = content.split(',', 1)[1]
            file_bytes = base64.b64decode(content)
            raw_text = _extract_docx_text(file_bytes)
            questions = _parse_word_questions(raw_text)
            results[filename] = {"count": len(questions), "questions": questions}
            total_found += len(questions)
        except Exception as e:
            results[filename] = {"count": 0, "questions": [], "error": str(e)}
    if total_found == 0:
        return jsonify({"status": "error", "message": "لم يتم العثور على أي أسئلة صالحة"}), 400
    return jsonify({"status": "success", "files": results, "total": total_found})


@app.route('/api/admin-save-quiz-word', methods=['POST'])
def api_admin_save_quiz_word():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    quiz_id = _clean_str(data.get('id'), 60)
    title = _clean_str(data.get('title', ''), 200)
    subject_id = _clean_str(data.get('subject_id', ''), 60)
    duration = int(data.get('duration') or 30)
    sections_config = data.get('sections', [])

    if not title or not subject_id:
        return jsonify({"status": "error", "message": "عنوان ومادة مطلوبين"}), 400
    if duration < 1 or duration > 300:
        return jsonify({"status": "error", "message": "المدة 1-300 دقيقة"}), 400
    if not sections_config or not isinstance(sections_config, list):
        return jsonify({"status": "error", "message": "لازم ملف واحد على الأقل"}), 400
    if not _quiz_is_teacher_of_subject(session['admin'], subject_id):
        return jsonify({"status": "error", "message": "غير مصرح بهذه المادة"}), 403

    subject = subjects_col.find_one({"id": subject_id})
    if not subject:
        return jsonify({"status": "error", "message": "المادة غير موجودة"}), 404

    sections_data = []
    for section in sections_config:
        section_name = _clean_str(section.get('name', ''), 100)
        all_qs = section.get('questions', [])
        num_pick = int(section.get('num_to_pick') or 0)
        mcq_pts = float(section.get('mcq_points') or 1)
        tf_pts = float(section.get('tf_points') or 1)
        if not all_qs or num_pick < 1:
            continue
        cleaned = []
        for q in all_qs:
            if not q.get('text'):
                continue
            cleaned.append({"text": q.get('text', ''), "type": q.get('type', 'mcq'),
                            "options": q.get('options', []), "correct": q.get('correct', 'a')})
        if not cleaned:
            continue
        actual_pick = min(num_pick, len(cleaned))
        sections_data.append({"name": section_name, "num_to_pick": actual_pick,
                              "mcq_points": mcq_pts, "tf_points": tf_pts,
                              "all_questions": cleaned})

    if not sections_data:
        return jsonify({"status": "error", "message": "لم يتم العثور على أسئلة صالحة"}), 400

    max_est = sum(s['num_to_pick'] * max(s['mcq_points'], s['tf_points']) for s in sections_data)
    min_est = sum(s['num_to_pick'] * min(s['mcq_points'], s['tf_points']) for s in sections_data)
    estimated_total = round((max_est + min_est) / 2, 2)

    flat_questions = []
    counter = 1
    for s in sections_data:
        for q in s['all_questions']:
            pts = s['tf_points'] if q['type'] == 'tf' else s['mcq_points']
            flat_questions.append({"id": f"q{counter}", "text": q['text'], "type": q['type'],
                                   "options": q['options'], "correct": q['correct'],
                                   "points": pts, "section": s['name']})
            counter += 1

    expected_count = _compute_expected_count({"sections_config": sections_data})

    if quiz_id:
        existing = quizzes_col.find_one({"id": quiz_id})
        if not existing or not _quiz_can_manage(session['admin'], existing):
            return jsonify({"status": "error", "message": "غير مصرح"}), 403
        quizzes_col.update_one({"id": quiz_id}, {"$set": {
            "title": title, "subject_id": subject_id,
            "subject_name": subject.get('name', ''),
            "duration": duration, "sections_config": sections_data,
            "questions": flat_questions, "total_points": estimated_total,
            "expected_questions": expected_count,
            "updated_at": _cairo_now()}})
    else:
        quiz_id = f"QUIZ_{int(time.time() * 1000)}"
        quizzes_col.insert_one({
            "id": quiz_id, "title": title, "subject_id": subject_id,
            "subject_name": subject.get('name', ''),
            "duration": duration, "sections_config": sections_data,
            "questions": flat_questions, "total_points": estimated_total,
            "expected_questions": expected_count,
            "is_open": True, "broadcast_active": False,
            "created_by": session['admin']['name'],
            "created_by_username": session['admin']['username'],
            "created_at": _cairo_now(), "created_at_dt": datetime.utcnow()})

    return jsonify({"status": "success", "quiz_id": quiz_id,
                    "questions_count": expected_count,
                    "total_points": estimated_total})


@app.route('/api/quiz-edit-question', methods=['POST'])
def api_quiz_edit_question():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    quiz_id = _clean_str(data.get('quiz_id'), 60)
    qid = _clean_str(data.get('qid'), 20)
    new_text = _clean_str(data.get('text', ''), 1000)
    new_options = data.get('options') or []
    new_correct = _clean_str(data.get('correct', 'a'), 5).lower()

    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403

    if quiz.get('sections_config') and qid.startswith('s') and '_q' in qid:
        try:
            parts = qid.split('_q')
            si = int(parts[0][1:]); qi = int(parts[1])
            if 0 <= si < len(quiz['sections_config']):
                section = quiz['sections_config'][si]
                all_qs = section.get('all_questions', [])
                if 0 <= qi < len(all_qs):
                    q = all_qs[qi]
                    if new_text: q['text'] = new_text
                    if new_options and len(new_options) >= 2: q['options'] = new_options[:4]
                    q['correct'] = new_correct
                    quizzes_col.update_one({"id": quiz_id}, {"$set": {"sections_config": quiz['sections_config']}})
                    return jsonify({"status": "success"})
        except Exception as e:
            return jsonify({"status": "error", "message": str(e)}), 500
        return jsonify({"status": "error", "message": "السؤال غير موجود"}), 404

    new_points = float(data.get('points') or 1)
    updated = False
    for q in quiz.get('questions', []):
        if q['id'] == qid:
            if new_text: q['text'] = new_text
            if new_options and len(new_options) >= 2: q['options'] = new_options[:4]
            q['correct'] = new_correct
            q['points'] = new_points
            updated = True
            break
    if not updated:
        return jsonify({"status": "error", "message": "السؤال غير موجود"}), 404
    total_points = round(sum(q.get('points', 1) for q in quiz['questions']), 2)
    quizzes_col.update_one({"id": quiz_id}, {"$set": {"questions": quiz['questions'], "total_points": total_points}})
    return jsonify({"status": "success", "total_points": total_points})


@app.route('/api/admin-edit-quiz-meta', methods=['POST'])
def api_admin_edit_quiz_meta():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    quiz_id = _clean_str(data.get('id'), 60)
    new_title = _clean_str(data.get('title', ''), 200)
    new_duration = int(data.get('duration') or 0)
    new_subject_id = _clean_str(data.get('subject_id', ''), 60)

    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403

    updates = {}
    if new_title:
        updates['title'] = new_title
    if 1 <= new_duration <= 300:
        updates['duration'] = new_duration
    if new_subject_id:
        if not _quiz_is_teacher_of_subject(session['admin'], new_subject_id):
            return jsonify({"status": "error", "message": "غير مصرح بالمادة"}), 403
        subj = subjects_col.find_one({"id": new_subject_id})
        if subj:
            updates['subject_id'] = new_subject_id
            updates['subject_name'] = subj.get('name', '')

    if not updates:
        return jsonify({"status": "error", "message": "لا يوجد تغييرات"}), 400

    updates['updated_at'] = _cairo_now()
    quizzes_col.update_one({"id": quiz_id}, {"$set": updates})
    return jsonify({"status": "success", "updates": updates})


@app.route('/api/admin-delete-quiz-section', methods=['POST'])
def api_admin_delete_quiz_section():
    if not _is_admin():
        return jsonify({"status": "unauthorized"}), 401
    data = request.get_json(force=True, silent=True) or {}
    quiz_id = _clean_str(data.get('quiz_id'), 60)
    section_name = _clean_str(data.get('section'), 100)
    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return jsonify({"status": "error", "message": "غير مصرح"}), 403

    if quiz.get('sections_config'):
        new_sections = [s for s in quiz['sections_config'] if s.get('name', '') != section_name]
        if len(new_sections) == len(quiz['sections_config']):
            return jsonify({"status": "error", "message": "القسم غير موجود"}), 404
        if not new_sections:
            return jsonify({"status": "error", "message": "لا يمكن حذف كل الأقسام"}), 400
        new_expected = _compute_expected_count({"sections_config": new_sections})
        quizzes_col.update_one({"id": quiz_id}, {"$set": {
            "sections_config": new_sections,
            "expected_questions": new_expected,
            "updated_at": _cairo_now()}})
        return jsonify({"status": "success",
                        "remaining": new_expected,
                        "total_points": quiz.get('total_points', 0)})

    remaining = [q for q in quiz.get('questions', []) if (q.get('section', '') or '') != section_name]
    if len(remaining) == len(quiz.get('questions', [])):
        return jsonify({"status": "error", "message": "القسم غير موجود"}), 404
    if not remaining:
        return jsonify({"status": "error", "message": "لا يمكن حذف كل الأقسام"}), 400
    for i, q in enumerate(remaining):
        q['id'] = f"q{i+1}"
    total_points = round(sum(q.get('points', 1) for q in remaining), 2)
    quizzes_col.update_one({"id": quiz_id}, {"$set": {
        "questions": remaining, "total_points": total_points,
        "expected_questions": len(remaining),
        "updated_at": _cairo_now()}})
    return jsonify({"status": "success", "remaining": len(remaining), "total_points": total_points})


@app.route('/api/quiz-excel-export')
def api_quiz_excel_export():
    if not _is_admin():
        return "Unauthorized", 401
    quiz_id = _clean_str(request.args.get('quiz_id'), 60)
    quiz = quizzes_col.find_one({"id": quiz_id})
    if not quiz or not _quiz_can_manage(session['admin'], quiz):
        return "Unauthorized", 403

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "openpyxl not installed", 500

    attempts = list(quiz_attempts_col.find({"quiz_id": quiz_id}).sort("started_at", -1))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Results"
    ws.sheet_view.rightToLeft = True

    ws.merge_cells('A1:I1')
    ws['A1'] = f"نتائج كويز: {quiz.get('title', '')}"
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="1a2136", end_color="1a2136", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:I2')
    ws['A2'] = (f"المادة: {quiz.get('subject_name', '')}  |  "
                f"عدد الأسئلة: {_compute_expected_count(quiz)}  |  "
                f"إجمالي الدرجات: {quiz.get('total_points', 0)}  |  "
                f"المدة: {quiz.get('duration', 0)} دقيقة")
    ws['A2'].font = Font(bold=True, size=11, color="1a2136")
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    headers = ['#', 'كود الطالب', 'اسم الطالب', 'الحالة', 'الدرجة', 'من',
               'البداية', 'التسليم', 'سبب الإلغاء']
    ws.append(headers)
    for cell in ws[3]:
        cell.fill = PatternFill(start_color="FFB300", end_color="FFB300", fill_type="solid")
        cell.font = Font(bold=True, color="000000", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 24

    for i, a in enumerate(attempts, 1):
        status_ar = {'completed': '✅ مكتمل', 'cancelled': '❌ ملغي',
                     'in_progress': '⏳ جاري'}.get(a.get('status'), a.get('status', ''))
        started = ''
        if a.get('started_at'):
            try:
                started = a['started_at'].strftime('%Y-%m-%d %H:%M')
            except Exception:
                started = str(a['started_at'])
        row = [i, a.get('student_id', ''), a.get('student_name', ''),
               status_ar,
               a.get('score', '') if a.get('status') == 'completed' else '',
               a.get('total_points', '') if a.get('status') == 'completed' else '',
               started, a.get('submitted_at', ''), a.get('cancelled_reason', '')]
        ws.append(row)
        last = ws.max_row
        if a.get('status') == 'completed':
            for cell in ws[last]:
                cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
        elif a.get('status') == 'cancelled':
            for cell in ws[last]:
                cell.fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

    widths = [6, 14, 28, 14, 10, 10, 18, 20, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A4'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe_title = re.sub(r'[^\w\u0600-\u06FF\s-]', '', quiz.get('title', 'Quiz'))[:30].strip().replace(' ', '_')
    filename = f"Quiz_{safe_title}_{int(time.time())}.xlsx"

    return Response(output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.route('/api/quiz-live-stats')
def api_quiz_live_stats():
    quiz_id = _clean_str(request.args.get('quiz_id'), 60)
    quiz = quizzes_col.find_one({"id": quiz_id},
        {"title": 1, "subject_name": 1, "questions": 1, "sections_config": 1,
         "total_points": 1, "duration": 1, "is_open": 1, "expected_questions": 1})
    if not quiz:
        return jsonify({"status": "error"}), 404
    completed = quiz_attempts_col.count_documents({"quiz_id": quiz_id, "status": "completed"})
    in_progress = quiz_attempts_col.count_documents({"quiz_id": quiz_id, "status": "in_progress"})
    return jsonify({"status": "success",
        "quiz": {"title": quiz.get('title', ''),
                 "subject_name": quiz.get('subject_name', ''),
                 "questions_count": quiz.get('expected_questions') or _compute_expected_count(quiz),
                 "total_points": quiz.get('total_points', 0),
                 "duration": quiz.get('duration', 0),
                 "is_open": quiz.get('is_open', False)},
        "completed": completed, "in_progress": in_progress,
        "total_attempts": completed + in_progress})


@app.route('/api/student-quizzes')
def api_student_quizzes():
    sid = _clean_str(request.args.get('student_id'), 20)
    if len(sid) != 7 or not sid.isdigit():
        return jsonify({"status": "error", "message": "ID غير صحيح"}), 400

    student = students_col.find_one({"student_id": sid})
    if not student:
        return jsonify({"status": "success", "quizzes": []})

    student_year = student.get('year') or ''
    student_dept = student.get('department') or 'عام (IT)'

    if not student_year:
        return jsonify({"status": "success", "quizzes": []})

    try:
        subs = list(subjects_col.find({
            "year": student_year,
            "department": {"$in": [student_dept, "عام (IT)"]}
        }, {"id": 1}))
        sub_ids = [s['id'] for s in subs]
    except Exception:
        sub_ids = []

    if not sub_ids:
        return jsonify({"status": "success", "quizzes": []})

    items = list(quizzes_col.find(
        {"subject_id": {"$in": sub_ids}, "is_open": True},
        {"questions.correct": 0, "_id": 0, "sections_config": 0}
    ).sort("created_at_dt", -1))

    for it in items:
        att = quiz_attempts_col.find_one(
            {"quiz_id": it['id'], "student_id": sid}, {"_id": 0})
        block = quiz_blocks_col.find_one(
            {"quiz_id": it['id'], "student_id": sid, "unblocked_at": None})

        it['my_status'] = 'blocked' if block else (
            att.get('status', 'not_started') if att else 'not_started'
        )
        it['questions_count'] = it.get('expected_questions') or _compute_expected_count(it)
        it['total_points'] = it.get('total_points', 0)
        it['broadcast_active'] = it.get('broadcast_active', False)
        it.pop('questions', None)

        if att and att.get('status') == 'completed':
            it['my_score'] = att.get('score')
            it['my_total'] = att.get('total_points')

    return jsonify({"status": "success", "quizzes": items})


@app.route('/api/quiz-precheck')
def api_quiz_precheck():
    qid = _clean_str(request.args.get('quiz_id'), 60)
    sid = _clean_str(request.args.get('student_id'), 20)
    if len(sid) != 7 or not sid.isdigit():
        return jsonify({"status": "error", "message": "ID غير صحيح"}), 400
    quiz = quizzes_col.find_one({"id": qid})
    if not quiz or not quiz.get('is_open'):
        return jsonify({"status": "error", "message": "غير متاح"}), 400

    student = students_col.find_one({"student_id": sid})
    if not student or not _student_can_access_quiz(student, quiz):
        return jsonify({"status": "error", "message": "غير متاح لمستواك"}), 403

    # ⚡ Auto-unblock transient disconnects
    block = quiz_blocks_col.find_one({"quiz_id": qid, "student_id": sid, "unblocked_at": None})
    if block:
        reason = (block.get('reason') or '').strip()
        blocked_dt = block.get('blocked_at_dt')
        if reason in ('انقطاع', 'Session ended from server', 'انقطاع شبكة') and blocked_dt:
            try:
                elapsed = (datetime.utcnow() - blocked_dt).total_seconds()
                if elapsed < 60:
                    quiz_blocks_col.delete_many({"quiz_id": qid, "student_id": sid})
                    quiz_attempts_col.delete_many({
                        "quiz_id": qid, "student_id": sid,
                        "status": "cancelled"
                    })
                    block = None
            except Exception:
                pass
        if block:
            return jsonify({"status": "blocked", "message": "محظور"})

    att = quiz_attempts_col.find_one({"quiz_id": qid, "student_id": sid}, {"_id": 0})
    if att:
        if att.get('status') == 'completed':
            return jsonify({"status": "already_done", "score": att.get('score'), "total": att.get('total_points')})
        if att.get('status') == 'in_progress':
            last_hb = att.get('last_heartbeat_dt') or att.get('started_at')
            if last_hb and (datetime.utcnow() - last_hb).total_seconds() > HEARTBEAT_GRACE_SEC:
                quiz_attempts_col.update_one({"attempt_id": att['attempt_id']},
                    {"$set": {"status": "cancelled", "cancelled_reason": "خرج من الصفحة", "cancelled_at": _cairo_now()}})
                quiz_blocks_col.update_one({"quiz_id": qid, "student_id": sid},
                    {"$setOnInsert": {"quiz_id": qid, "student_id": sid,
                                      "student_name": att.get('student_name', ''),
                                      "reason": "انقطاع", "blocked_at": _cairo_now(),
                                      "blocked_at_dt": datetime.utcnow(),
                                      "unblocked_at": None}}, upsert=True)
                return jsonify({"status": "blocked", "message": "خرجت — تم الحظر"})
            return jsonify({"status": "can_resume"})

    return jsonify({"status": "can_start", "title": quiz.get('title'),
                    "duration": quiz.get('duration'),
                    "questions_count": _compute_expected_count(quiz)})


@app.route('/api/start-quiz', methods=['POST'])
def api_start_quiz():
    data = request.get_json(force=True, silent=True) or {}
    qid = _clean_str(data.get('quiz_id'), 60)
    sid = _clean_str(data.get('student_id'), 20)
    sname = _clean_str(data.get('student_name'), 100)
    if len(sid) != 7 or not sid.isdigit():
        return jsonify({"status": "error", "message": "ID غير صحيح"}), 400
    quiz = quizzes_col.find_one({"id": qid})
    if not quiz or not quiz.get('is_open'):
        return jsonify({"status": "error", "message": "غير متاح"}), 400

    student = students_col.find_one({"student_id": sid})
    if not student or not _student_can_access_quiz(student, quiz):
        return jsonify({"status": "error", "message": "غير متاح لمستواك"}), 403

    if quiz_blocks_col.find_one({"quiz_id": qid, "student_id": sid, "unblocked_at": None}):
        return jsonify({"status": "blocked", "message": "محظور"})

    existing = quiz_attempts_col.find_one({"quiz_id": qid, "student_id": sid})
    if existing:
        if existing.get('status') == 'completed':
            return jsonify({"status": "already_done"})
        last_hb = existing.get('last_heartbeat_dt') or existing.get('started_at')
        if last_hb and (datetime.utcnow() - last_hb).total_seconds() > HEARTBEAT_GRACE_SEC:
            quiz_attempts_col.update_one({"attempt_id": existing['attempt_id']},
                {"$set": {"status": "cancelled", "cancelled_reason": "انقطاع", "cancelled_at": _cairo_now()}})
            quiz_blocks_col.update_one({"quiz_id": qid, "student_id": sid},
                {"$setOnInsert": {"quiz_id": qid, "student_id": sid,
                                  "student_name": sname, "reason": "انقطاع",
                                  "blocked_at": _cairo_now(),
                                  "blocked_at_dt": datetime.utcnow(),
                                  "unblocked_at": None}}, upsert=True)
            return jsonify({"status": "blocked", "message": "تم الحظر"})
        if existing.get('generated_questions'):
            generated_qs = existing['generated_questions']
            total_pts = existing.get('total_points', 0)
        else:
            generated_qs = _generate_attempt_questions(quiz)
            total_pts = round(sum(q.get('points', 1) for q in generated_qs), 2)
            quiz_attempts_col.update_one({"attempt_id": existing['attempt_id']},
                {"$set": {"generated_questions": generated_qs, "total_points": total_pts}})
        attempt = existing
    else:
        generated_qs = _generate_attempt_questions(quiz)
        total_pts = round(sum(q.get('points', 1) for q in generated_qs), 2)
        attempt_id = f"AT_{int(time.time()*1000)}_{secrets.token_hex(3)}"
        now = datetime.utcnow()
        quiz_attempts_col.insert_one({
            "attempt_id": attempt_id, "quiz_id": qid, "student_id": sid,
            "student_name": sname, "started_at": now, "last_heartbeat_dt": now,
            "expires_at": now + timedelta(minutes=quiz.get('duration', 30)),
            "status": "in_progress", "answers": {}, "score": None,
            "generated_questions": generated_qs, "total_points": total_pts, "violations": 0})
        attempt = quiz_attempts_col.find_one({"attempt_id": attempt_id})

    public_qs = [{"id": q['id'], "text": q['text'], "type": q['type'],
                  "options": q['options'], "points": q['points']} for q in generated_qs]
    expires_in = max(0, int((attempt['expires_at'] - datetime.utcnow()).total_seconds()))
    return jsonify({"status": "success", "attempt_id": attempt['attempt_id'],
        "quiz": {"id": quiz['id'], "title": quiz['title'],
                 "duration": quiz['duration'], "questions": public_qs,
                 "total_points": total_pts},
        "expires_in": expires_in, "server_now": datetime.utcnow().isoformat() + "Z"})


# ============================================================
#  💓 HEARTBEAT — saves answers + health check
# ============================================================
@app.route('/api/quiz-heartbeat', methods=['POST'])
def api_quiz_heartbeat():
    data = request.get_json(force=True, silent=True) or {}
    att_id = _clean_str(data.get('attempt_id'), 60)
    sid = _clean_str(data.get('student_id'), 20)
    answers = data.get('answers')

    update_doc = {"$set": {"last_heartbeat_dt": datetime.utcnow()}}
    if isinstance(answers, dict):
        update_doc["$set"]["answers"] = answers

    r = quiz_attempts_col.update_one(
        {"attempt_id": att_id, "student_id": sid, "status": "in_progress"},
        update_doc)

    if r.matched_count:
        return jsonify({"status": "ok"})

    # Already finished?
    done = quiz_attempts_col.find_one(
        {"attempt_id": att_id, "student_id": sid},
        {"_id": 0, "status": 1, "score": 1, "total_points": 1})
    if done and done.get('status') == 'completed':
        return jsonify({"status": "completed",
                        "score": done.get('score'),
                        "total": done.get('total_points')})
    return jsonify({"status": "expired"})


# ============================================================
#  ⏱️ TIME SYNC — auto-grade on expiry, no more "0 score + block"
# ============================================================
@app.route('/api/quiz-time-sync')
def api_quiz_time_sync():
    att_id = _clean_str(request.args.get('attempt_id'), 60)
    sid = _clean_str(request.args.get('student_id'), 20)

    attempt = quiz_attempts_col.find_one({"attempt_id": att_id, "student_id": sid})
    if not attempt:
        return jsonify({"status": "not_active"})

    if attempt.get('status') == 'completed':
        return jsonify({"status": "expired", "time_left": 0,
                        "score": attempt.get('score'),
                        "total": attempt.get('total_points'),
                        "auto_submitted": attempt.get('auto_submitted', False)})

    if attempt.get('status') != 'in_progress':
        return jsonify({"status": "not_active"})

    now = datetime.utcnow()
    time_left = (attempt['expires_at'] - now).total_seconds()

    if time_left <= 0:
        quiz = quizzes_col.find_one({"id": attempt['quiz_id']})
        if quiz:
            _auto_grade_attempt(attempt, quiz)
            updated = quiz_attempts_col.find_one({"attempt_id": att_id})
            return jsonify({"status": "expired", "time_left": 0,
                            "score": updated.get('score'),
                            "total": updated.get('total_points'),
                            "auto_submitted": True})
        return jsonify({"status": "expired", "time_left": 0})

    return jsonify({"status": "active", "time_left": int(time_left),
                    "server_now": now.isoformat() + "Z"})


# ============================================================
#  🚨 VIOLATION — respects expiry (no more false "cheat" after time up)
# ============================================================
@app.route('/api/quiz-violation', methods=['POST'])
def api_quiz_violation():
    data = request.get_json(force=True, silent=True) or {}
    att_id = _clean_str(data.get('attempt_id'), 60)
    sid = _clean_str(data.get('student_id'), 20)
    reason = _clean_str(data.get('reason', 'unknown'), 200)

    attempt = quiz_attempts_col.find_one(
        {"attempt_id": att_id, "student_id": sid, "status": "in_progress"})
    if not attempt:
        return jsonify({"status": "ignored"})

    # ✅ Time is up → auto-grade, DO NOT block
    if datetime.utcnow() >= attempt['expires_at']:
        quiz = quizzes_col.find_one({"id": attempt['quiz_id']})
        if quiz:
            _auto_grade_attempt(attempt, quiz)
        return jsonify({"status": "expired"})

    # Real violation → block
    quiz_attempts_col.update_one({"attempt_id": att_id},
        {"$set": {"status": "cancelled", "cancelled_reason": reason, "cancelled_at": _cairo_now()},
         "$inc": {"violations": 1}})
    quiz_blocks_col.update_one({"quiz_id": attempt['quiz_id'], "student_id": sid},
        {"$setOnInsert": {"quiz_id": attempt['quiz_id'], "student_id": sid,
                          "student_name": attempt.get('student_name', ''),
                          "reason": reason, "blocked_at": _cairo_now(),
                          "blocked_at_dt": datetime.utcnow(),
                          "unblocked_at": None}}, upsert=True)
    return jsonify({"status": "blocked"})


# ============================================================
#  ✅ SUBMIT — accepts late submission + returns existing result
# ============================================================
@app.route('/api/submit-quiz', methods=['POST'])
def api_submit_quiz():
    data = request.get_json(force=True, silent=True) or {}
    att_id = _clean_str(data.get('attempt_id'), 60)
    sid = _clean_str(data.get('student_id'), 20)
    answers = data.get('answers') or {}
    if not isinstance(answers, dict):
        return jsonify({"status": "error", "message": "إجابات غير صالحة"}), 400

    attempt = quiz_attempts_col.find_one({"attempt_id": att_id, "student_id": sid})
    if not attempt:
        return jsonify({"status": "error", "message": "غير موجودة"}), 404

    # ⚡ Already completed (from time-sync or prior submit) — return result
    if attempt.get('status') == 'completed':
        return jsonify({"status": "success",
                        "score": attempt.get('score'),
                        "total": attempt.get('total_points'),
                        "details": attempt.get('details', []),
                        "auto_submitted": attempt.get('auto_submitted', False)})

    if attempt.get('status') != 'in_progress':
        return jsonify({"status": "error", "message": "منتهية"}), 400

    quiz = quizzes_col.find_one({"id": attempt['quiz_id']})
    if not quiz:
        return jsonify({"status": "error", "message": "الكويز محذوف"}), 404

    # Save latest answers first (so auto-grade uses them if needed)
    quiz_attempts_col.update_one(
        {"attempt_id": att_id, "status": "in_progress"},
        {"$set": {"answers": answers}})

    # If late, still accept — mark auto_submitted
    if datetime.utcnow() > attempt['expires_at']:
        attempt = quiz_attempts_col.find_one({"attempt_id": att_id})
        _auto_grade_attempt(attempt, quiz)
        updated = quiz_attempts_col.find_one({"attempt_id": att_id})
        return jsonify({"status": "success",
                        "score": updated.get('score'),
                        "total": updated.get('total_points'),
                        "details": updated.get('details', []),
                        "auto_submitted": True})

    # Normal grading
    attempt_questions = attempt.get('generated_questions') or quiz.get('questions', [])
    score = 0.0; total = 0.0; details = []
    for q in attempt_questions:
        total += float(q.get('points', 1))
        ans = answers.get(q['id'])
        ok = (ans == q['correct'])
        if ok:
            score += float(q.get('points', 1))
        details.append({"qid": q['id'], "text": q['text'], "user": ans,
                        "correct": q['correct'], "ok": ok, "points": q.get('points', 1)})

    r = quiz_attempts_col.find_one_and_update(
        {"attempt_id": att_id, "status": "in_progress"},
        {"$set": {"status": "completed", "score": round(score, 2),
                  "total_points": round(total, 2), "answers": answers,
                  "details": details, "submitted_at": _cairo_now(),
                  "submitted_at_dt": datetime.utcnow()}},
        return_document=True)
    if not r:
        return jsonify({"status": "error", "message": "تم التسليم مسبقاً"}), 400
    return jsonify({"status": "success", "score": round(score, 2),
                    "total": round(total, 2), "details": details})


@app.route('/api/export-attendance-csv')
def export_attendance_csv():
    if not _is_admin():
        return "Unauthorized", 401
    role = session['admin']['role']
    query = {}
    sess_id = _clean_str(request.args.get('session_id'), 64)
    raw_filename = "Attendance_Report"
    if sess_id:
        query["session_id"] = sess_id
        sess = sessions_col.find_one({"session_id": sess_id})
        if sess:
            if role == 'ta':
                curr = users_col.find_one({"username": session['admin']['username']})
                if sess.get('created_by_username') != curr.get('username'):
                    return "Unauthorized", 403
            elif role == 'doctor':
                curr = users_col.find_one({"username": session['admin']['username']})
                if sess.get('subject_id') not in (curr.get('allowed_subjects') or []):
                    return "Unauthorized", 403
            raw_filename = f"{sess.get('subject_name','')}_{sess.get('title','')}"
    records = list(attendance_col.find(query, {"_id": 0}).sort("timestamp", -1).limit(20000))
    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output)
    writer.writerow(['كود الطالب', 'اسم الطالب', 'الفرقة', 'القسم', 'المادة',
                     'نوع الجلسة', 'عنوان الجلسة', 'توقيت الحضور', 'طريقة التسجيل', 'IP'])
    for r in records:
        scan_type = r.get('scan_type', '')
        if r.get('is_manual') or 'ADMIN' in str(r.get('ip', '')):
            method = 'يدوي (آدمن)'
        elif scan_type == 'qr':
            method = 'QR تلقائي'
        else:
            method = 'كود يدوي'
        writer.writerow([r.get('student_id', ''), r.get('student_name', ''),
                         r.get('year', ''), r.get('department', ''),
                         r.get('subject_name', ''), r.get('session_type', ''),
                         r.get('session_title', ''), r.get('timestamp', ''),
                         method, r.get('ip', '')])
    csv_data = output.getvalue()
    safe_ascii_name = f"Attendance_{int(time.time())}.csv"
    clean_name = "".join(ch for ch in raw_filename if ch.isalnum() or ch in " _-")[:80] or "Report"
    encoded_utf8_name = urllib.parse.quote(f"{clean_name}_{int(time.time())}.csv")
    disposition_header = f"attachment; filename=\"{safe_ascii_name}\"; filename*=UTF-8''{encoded_utf8_name}"
    return Response(csv_data, mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": disposition_header})


if __name__ == '__main__':
    debug_mode = os.environ.get("FLASK_DEBUG", "0") == "1"
    print("=" * 65)
    print("🚀 Nexus Attendance Server Running")
    print("🔑 Admin: http://127.0.0.1:8080/secure-auth-gateway-2026-x9v2-pl7q-a84m")
    print("🖥️  Display: http://127.0.0.1:8080/dis")
    print(f"🔧 Debug mode: {'ON' if debug_mode else 'OFF'}")
    print("=" * 65)
    app.run(host='0.0.0.0', port=8080, debug=debug_mode, use_reloader=False)
